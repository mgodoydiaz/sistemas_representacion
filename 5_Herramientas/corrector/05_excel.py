# -*- coding: utf-8 -*-
"""
Genera la planilla de notas de la Actividad 1.
PCI1119 Sistemas de Representacion.

Cruza tres fuentes:
  - Nomina oficial (29 estudiantes)
  - salida/inventario.csv        (formato en que entrego cada uno)
  - salida/lote_resultados.csv   (correccion automatica, donde fue confiable)

Salida: ../Notas_Actividad1_<codigo de la seccion>.xlsx
"""

import csv
import os
import sys
from collections import defaultdict

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import contexto
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
CARAS = {1: 30, 2: 44, 3: 43, 4: 48, 5: 47}
TOTAL_CARAS = sum(CARAS.values())

DESCUENTOS = [
    ("PDF", 0, "Formato solicitado en el enunciado"),
    ("Word", -10, "Entrego .docx en vez de PDF"),
    ("Varios PDF", -10, "Entrego varios PDF sueltos en vez de un solo archivo"),
    ("Imagen", -20, "Entrego JPG, PNG, screenshot o foto en vez de PDF"),
    ("Sin entrega", 0, "No entrego; el puntaje ya queda en cero por si solo"),
]

MAPA_FORMATO = {
    "pdf": "PDF",
    "pdf_multiple": "Varios PDF",
    "docx": "Word",
    "applet_jpg": "Imagen",
    "screenshot": "Imagen",
    "foto": "Imagen",
    "mixto": "Imagen",
    "sin_archivos": "Sin entrega",
}

# Tipografia y colores
FUENTE = "Arial"
AZUL = "2E5E8C"
GRIS_CAB = "D9D9D9"
AMARILLO = "FFF2CC"   # celdas que Miguel debe llenar
VERDE = "C6EFCE"
AMBAR = "FFEB9C"
ROJO = "FFC7CE"
BORDE = Border(*[Side(style="thin", color="BFBFBF")] * 4)


def leer_csv(ruta, sep=","):
    with open(ruta, "r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f, delimiter=sep))


def cargar_nomina():
    ruta = contexto.ruta_nomina()
    wb = openpyxl.load_workbook(ruta, data_only=True)
    ws = contexto.hoja_nomina(wb)
    filas = []
    for r in ws.iter_rows(min_row=6, values_only=True):
        if not r or not r[0]:
            continue
        filas.append({"n": r[0], "apellidos": r[1], "nombres": r[2],
                      "rut": r[4], "correo": r[5]})
    return filas


def main():
    nomina = cargar_nomina()
    # El inventario puede no existir todavia: en una seccion recien montada la
    # planilla se genera en blanco, antes de que haya entregas que extraer.
    ruta_inv = os.path.join(contexto.raiz(), "salida", "inventario.csv")
    hay_inventario = os.path.exists(ruta_inv)
    inventario = ({f["usuario"]: f for f in leer_csv(ruta_inv)}
                  if hay_inventario else {})
    if not hay_inventario:
        print("Aviso: no hay salida/inventario.csv todavia.")
        print("Genero la planilla en blanco, con el formato de entrega en PDF por defecto.")
        print("Cuando corras 01_extraer.py y vuelvas a generarla, se llenara sola.")

    # Correccion automatica agregada por alumno y lamina
    auto = defaultdict(lambda: defaultdict(int))
    vistas = defaultdict(set)
    ruta_lote = os.path.join(contexto.raiz(), "salida", "lote_resultados.csv")
    if os.path.exists(ruta_lote):
        for f in leer_csv(ruta_lote, sep=";"):
            lam = int(f["lamina"])
            vistas[f["usuario"]].add(lam)
            if f["correcto"] == "1":
                auto[f["usuario"]][lam] += 1

    # Resultados de 08_corregir_curso.py, que corrige alumno por alumno sobre
    # los archivos crudos. Tiene prioridad sobre lote_resumen.csv porque cubre
    # tambien las entregas que la normalizacion automatica no pudo recortar.
    curso = {}
    ruta_curso = os.path.join(contexto.raiz(), "salida", "curso_resultados.csv")
    if os.path.exists(ruta_curso):
        for f in leer_csv(ruta_curso, sep=";"):
            curso[f["usuario"]] = f

    resumen_lote = {}
    ruta_res = os.path.join(contexto.raiz(), "salida", "lote_resumen.csv")
    if os.path.exists(ruta_res):
        for f in leer_csv(ruta_res, sep=";"):
            resumen_lote[f["usuario"]] = f

    wb = openpyxl.Workbook()

    # ------------------------------------------------------------ Parametros
    wp = wb.active
    wp.title = "Parametros"
    wp["A1"] = "PARAMETROS DE CORRECCION - ACTIVIDAD 1"
    wp["A1"].font = Font(name=FUENTE, size=13, bold=True, color=AZUL)
    wp["A2"] = "Cambiar un valor aqui recalcula toda la hoja Notas."
    wp["A2"].font = Font(name=FUENTE, size=9, italic=True, color="666666")

    wp["A4"] = "Caras por lamina"
    wp["A4"].font = Font(name=FUENTE, bold=True)
    for i, (lam, n) in enumerate(sorted(CARAS.items())):
        wp.cell(row=5 + i, column=1, value="Lamina %d" % lam).font = Font(name=FUENTE)
        c = wp.cell(row=5 + i, column=2, value=n)
        c.font = Font(name=FUENTE, color="0000FF")
    wp["A10"] = "Total de caras evaluadas"
    wp["A10"].font = Font(name=FUENTE, bold=True)
    wp["B10"] = "=SUM(B5:B9)"
    wp["B10"].font = Font(name=FUENTE, bold=True)

    wp["A12"] = "Descuento por formato de entrega"
    wp["A12"].font = Font(name=FUENTE, bold=True)
    wp["A13"] = "Formato"; wp["B13"] = "Descuento"; wp["C13"] = "Criterio"
    for c in ("A13", "B13", "C13"):
        wp[c].font = Font(name=FUENTE, bold=True)
        wp[c].fill = PatternFill("solid", fgColor=GRIS_CAB)
    for i, (nombre, desc, crit) in enumerate(DESCUENTOS):
        wp.cell(row=14 + i, column=1, value=nombre).font = Font(name=FUENTE)
        cd = wp.cell(row=14 + i, column=2, value=desc)
        cd.font = Font(name=FUENTE, color="0000FF")
        cd.fill = PatternFill("solid", fgColor=AMARILLO)
        wp.cell(row=14 + i, column=3, value=crit).font = Font(name=FUENTE, size=9)

    wp["A19"] = "Fuente de los descuentos: definidos por el profesor. El enunciado pedia entrega en PDF."
    wp["A19"].font = Font(name=FUENTE, size=9, italic=True, color="666666")
    wp["A20"] = "Fuente del conteo de caras: segmentacion automatica de las 5 pautas del applet educacionplastica.net"
    wp["A20"].font = Font(name=FUENTE, size=9, italic=True, color="666666")
    wp["A21"] = "Puntuacion binaria: cada cara vale 1 si el color calza con la pauta, 0 en cualquier otro caso."
    wp["A21"].font = Font(name=FUENTE, size=9, italic=True, color="666666")
    wp.column_dimensions["A"].width = 30
    wp.column_dimensions["B"].width = 12
    wp.column_dimensions["C"].width = 60

    # ----------------------------------------------------------------- Notas
    ws = wb.create_sheet("Notas")
    ws["A1"] = "ACTIVIDAD 1 - RECONOCIMIENTO DE VISTAS CON COLOR"
    ws["A1"].font = Font(name=FUENTE, size=14, bold=True, color=AZUL)
    ws["A2"] = contexto.encabezado()
    ws["A2"].font = Font(name=FUENTE, size=10, color="666666")
    ws["A3"] = ("Llena SOLO las celdas amarillas (caras correctas por lamina). El resto se calcula solo. "
                "Usa el corrector: corrector\\Corregir.bat")
    ws["A3"].font = Font(name=FUENTE, size=9, italic=True, color="806000")
    ws["A3"].fill = PatternFill("solid", fgColor=AMARILLO)

    cabeceras = ["N", "Apellidos", "Nombres", "Correo institucional",
                 "Formato de entrega", "Descuento formato",
                 "L1", "L2", "L3", "L4", "L5",
                 "Caras correctas", "Caras posibles", "Puntaje base",
                 "NOTA 0-100", "Estado", "Observaciones"]
    FIL_CAB = 5
    for j, h in enumerate(cabeceras, start=1):
        c = ws.cell(row=FIL_CAB, column=j, value=h)
        c.font = Font(name=FUENTE, bold=True, size=10,
                      color="FFFFFF" if j == 15 else "000000")
        c.fill = PatternFill("solid", fgColor=AZUL if j == 15 else GRIS_CAB)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDE

    # fila de maximos bajo L1..L5
    ws.cell(row=FIL_CAB + 1, column=6, value="maximo:").font = Font(
        name=FUENTE, size=8, italic=True, color="666666")
    ws.cell(row=FIL_CAB + 1, column=6).alignment = Alignment(horizontal="right")
    for i, lam in enumerate(sorted(CARAS)):
        c = ws.cell(row=FIL_CAB + 1, column=7 + i, value="=Parametros!$B$%d" % (5 + i))
        c.font = Font(name=FUENTE, size=8, italic=True, color="666666")
        c.alignment = Alignment(horizontal="center")

    fila = FIL_CAB + 2
    primera = fila
    for est in nomina:
        usuario = (est["correo"] or "").split("@")[0]
        inv = inventario.get(usuario, {})
        if hay_inventario:
            tipo = inv.get("tipo_entrega", "sin_archivos")
            formato = MAPA_FORMATO.get(tipo, "Imagen")
        else:
            # Sin inventario no se sabe como entrego nadie. Dejarlo en
            # "Sin entrega" pondria a todo el curso en cero, asi que se asume
            # el formato pedido y Miguel lo ajusta alumno por alumno.
            formato = "PDF"

        obs = []
        if inv.get("observacion"):
            obs.append(inv["observacion"])

        # prellenado automatico solo donde la correccion fue confiable
        prellenar = {}
        res = resumen_lote.get(usuario)
        if formato == "Sin entrega":
            prellenar = {l: 0 for l in CARAS}
        elif usuario in curso:
            c = curso[usuario]
            faltan_c = [x for x in (c.get("laminas_faltantes") or "").split(";") if x]
            for lam in sorted(CARAS):
                val = c.get("L%d" % lam, "")
                if val not in ("", None):
                    prellenar[lam] = int(float(val))
            if not prellenar:
                obs.append("El corrector no reconocio ninguna lamina. "
                           "Puede ser otro ejercicio o un recorte imposible. Revisar a mano")
            else:
                obs.append("Prellenado por el corrector automatico, verificar")
                if faltan_c:
                    obs.append("No se encontraron las laminas %s" % ", ".join(faltan_c))
                try:
                    if float(c.get("confianza_minima") or 1) < 0.85:
                        obs.append("Confianza de lectura %s, conviene revisar"
                                   % c.get("confianza_minima"))
                except ValueError:
                    pass
        elif res and res.get("revisar") == "no":
            prellenar = {l: auto[usuario].get(l, 0) for l in sorted(vistas[usuario])}
            obs.append("Prellenado por el corrector automatico, verificar")
        elif usuario == "gabriela.sepulveda2026":
            prellenar = {l: auto[usuario].get(l, 0) for l in sorted(vistas[usuario])}
            obs.append("Laminas 2 a 5 prellenadas por el corrector; falta la lamina 1")

        ws.cell(row=fila, column=1, value=est["n"]).font = Font(name=FUENTE, size=10)
        ws.cell(row=fila, column=2, value=est["apellidos"]).font = Font(name=FUENTE, size=10)
        ws.cell(row=fila, column=3, value=est["nombres"]).font = Font(name=FUENTE, size=10)
        ws.cell(row=fila, column=4, value=est["correo"]).font = Font(name=FUENTE, size=9)
        cf = ws.cell(row=fila, column=5, value=formato)
        cf.font = Font(name=FUENTE, size=10, color="0000FF")
        cf.fill = PatternFill("solid", fgColor=AMARILLO)
        cf.alignment = Alignment(horizontal="center")

        ws.cell(row=fila, column=6,
                value="=IFERROR(INDEX(Parametros!$B$14:$B$18,MATCH($E%d,Parametros!$A$14:$A$18,0)),0)" % fila
                ).font = Font(name=FUENTE, size=10)

        for i, lam in enumerate(sorted(CARAS)):
            c = ws.cell(row=fila, column=7 + i)
            if lam in prellenar:
                c.value = prellenar[lam]
            c.font = Font(name=FUENTE, size=10, color="0000FF")
            c.fill = PatternFill("solid", fgColor=AMARILLO)
            c.alignment = Alignment(horizontal="center")
            c.border = BORDE

        ws.cell(row=fila, column=12, value="=SUM($G%d:$K%d)" % (fila, fila)
                ).font = Font(name=FUENTE, size=10)
        ws.cell(row=fila, column=13, value="=Parametros!$B$10"
                ).font = Font(name=FUENTE, size=10, color="008000")
        ws.cell(row=fila, column=14,
                value="=IF($M%d=0,0,100*$L%d/$M%d)" % (fila, fila, fila))
        ws.cell(row=fila, column=14).font = Font(name=FUENTE, size=10)
        ws.cell(row=fila, column=14).number_format = "0.0"

        cn = ws.cell(row=fila, column=15,
                     value="=MAX(0,ROUND($N%d+$F%d,1))" % (fila, fila))
        cn.font = Font(name=FUENTE, size=11, bold=True)
        cn.number_format = "0.0"
        cn.alignment = Alignment(horizontal="center")
        cn.border = BORDE

        ws.cell(row=fila, column=16,
                value=('=IF($E{0}="Sin entrega","Sin entrega",'
                       'IF(COUNT($G{0}:$K{0})<5,"Pendiente","Listo"))').format(fila)
                ).font = Font(name=FUENTE, size=9)
        ws.cell(row=fila, column=16).alignment = Alignment(horizontal="center")

        ws.cell(row=fila, column=17, value=" | ".join(obs)
                ).font = Font(name=FUENTE, size=9, color="666666")
        fila += 1

    ultima = fila - 1

    # ------------------------------------------------------------- resumen
    fr = ultima + 2
    ws.cell(row=fr, column=2, value="Resumen del curso").font = Font(
        name=FUENTE, bold=True, color=AZUL)
    etiquetas = [
        ("Entregaron", '=COUNTIF($E{a}:$E{b},"<>Sin entrega")'),
        ("No entregaron", '=COUNTIF($E{a}:$E{b},"Sin entrega")'),
        ("Aun pendientes de corregir", '=COUNTIF($P{a}:$P{b},"Pendiente")'),
        ("Promedio del curso", "=ROUND(AVERAGE($O{a}:$O{b}),1)"),
        ("Promedio de quienes entregaron",
         '=IFERROR(ROUND(SUMIF($E{a}:$E{b},"<>Sin entrega",$O{a}:$O{b})/'
         'COUNTIF($E{a}:$E{b},"<>Sin entrega"),1),0)'),
        ("Aprobados (60 o mas)", '=COUNTIF($O{a}:$O{b},">=60")'),
        ("Nota mas alta", "=MAX($O{a}:$O{b})"),
        ("Nota mas baja", "=MIN($O{a}:$O{b})"),
    ]
    for i, (et, fo) in enumerate(etiquetas):
        ws.cell(row=fr + 1 + i, column=2, value=et).font = Font(name=FUENTE, size=10)
        c = ws.cell(row=fr + 1 + i, column=4,
                    value=fo.format(a=primera, b=ultima))
        c.font = Font(name=FUENTE, size=10, bold=True)
        c.number_format = "0.0"

    # --------------------------------------------------- validacion y formato
    for i, lam in enumerate(sorted(CARAS)):
        col = get_column_letter(7 + i)
        dv = DataValidation(type="whole", operator="between",
                            formula1="0", formula2=str(CARAS[lam]),
                            allow_blank=True, showErrorMessage=True)
        dv.errorTitle = "Valor fuera de rango"
        dv.error = "La lamina %d tiene %d caras. Ingresa un entero entre 0 y %d." % (
            lam, CARAS[lam], CARAS[lam])
        ws.add_data_validation(dv)
        dv.add("%s%d:%s%d" % (col, primera, col, ultima))

    dvf = DataValidation(type="list",
                         formula1='"PDF,Varios PDF,Word,Imagen,Sin entrega"',
                         allow_blank=False, showErrorMessage=True)
    dvf.errorTitle = "Formato no valido"
    dvf.error = "Elige PDF, Varios PDF, Word, Imagen o Sin entrega."
    ws.add_data_validation(dvf)
    dvf.add("E%d:E%d" % (primera, ultima))

    rango = "O%d:O%d" % (primera, ultima)
    ws.conditional_formatting.add(rango, CellIsRule(
        operator="greaterThanOrEqual", formula=["60"],
        fill=PatternFill("solid", bgColor=VERDE)))
    ws.conditional_formatting.add(rango, CellIsRule(
        operator="between", formula=["40", "59.999"],
        fill=PatternFill("solid", bgColor=AMBAR)))
    ws.conditional_formatting.add(rango, CellIsRule(
        operator="lessThan", formula=["40"],
        fill=PatternFill("solid", bgColor=ROJO)))
    ws.conditional_formatting.add("P%d:P%d" % (primera, ultima), CellIsRule(
        operator="equal", formula=['"Pendiente"'],
        fill=PatternFill("solid", bgColor=AMBAR)))

    anchos = {"A": 5, "B": 24, "C": 22, "D": 32, "E": 15, "F": 12,
              "G": 6, "H": 6, "I": 6, "J": 6, "K": 6, "L": 14, "M": 13,
              "N": 12, "O": 12, "P": 12, "Q": 52}
    for col, w in anchos.items():
        ws.column_dimensions[col].width = w
    ws.row_dimensions[FIL_CAB].height = 34
    ws.freeze_panes = "E%d" % primera
    ws.auto_filter.ref = "A%d:Q%d" % (FIL_CAB, ultima)

    # ------------------------------------------------- Estadistica por cara
    we = wb.create_sheet("Estadistica caras")
    we["A1"] = "ERRORES POR CARA - que caras confunden mas al curso"
    we["A1"].font = Font(name=FUENTE, size=13, bold=True, color=AZUL)
    we["A2"] = ("Calculado sobre los alumnos que el corrector automatico pudo procesar. "
                "Es una foto parcial, no del curso completo.")
    we["A2"].font = Font(name=FUENTE, size=9, italic=True, color="666666")

    cab = ["Lamina", "Pieza", "Cara", "Color correcto", "Alumnos revisados",
           "Errores", "% de error", "Colores equivocados mas frecuentes"]
    for j, h in enumerate(cab, start=1):
        c = we.cell(row=4, column=j, value=h)
        c.font = Font(name=FUENTE, bold=True, size=10)
        c.fill = PatternFill("solid", fgColor=GRIS_CAB)
        c.alignment = Alignment(horizontal="center", wrap_text=True)

    agregado = {}
    if os.path.exists(ruta_lote):
        for f in leer_csv(ruta_lote, sep=";"):
            k = (int(f["lamina"]), int(f["pieza"]), f["cara_id"])
            d = agregado.setdefault(k, {"esperado": f["color_esperado"], "n": 0,
                                        "err": 0, "puestos": defaultdict(int)})
            d["n"] += 1
            if f["correcto"] != "1":
                d["err"] += 1
                d["puestos"][f["color_detectado"]] += 1

    r = 5
    for k in sorted(agregado, key=lambda x: (x[0], x[1], x[2])):
        d = agregado[k]
        puestos = sorted(d["puestos"].items(), key=lambda t: -t[1])
        we.cell(row=r, column=1, value=k[0])
        we.cell(row=r, column=2, value=k[1])
        we.cell(row=r, column=3, value=k[2])
        we.cell(row=r, column=4, value=d["esperado"])
        we.cell(row=r, column=5, value=d["n"])
        we.cell(row=r, column=6, value=d["err"])
        ce = we.cell(row=r, column=7, value="=IF($E%d=0,0,$F%d/$E%d)" % (r, r, r))
        ce.number_format = "0%"
        we.cell(row=r, column=8,
                value=", ".join("%s (%d)" % (c, n) for c, n in puestos[:3]))
        for j in range(1, 9):
            we.cell(row=r, column=j).font = Font(name=FUENTE, size=10)
        r += 1
    if r > 5:
        we.auto_filter.ref = "A4:H%d" % (r - 1)
        we.conditional_formatting.add("G5:G%d" % (r - 1), CellIsRule(
            operator="greaterThanOrEqual", formula=["0.3"],
            fill=PatternFill("solid", bgColor=ROJO)))
    for col, w in {"A": 8, "B": 8, "C": 10, "D": 15, "E": 17, "F": 10,
                   "G": 11, "H": 42}.items():
        we.column_dimensions[col].width = w
    we.freeze_panes = "A5"

    salida = os.path.join(contexto.seccion_dir(), contexto.nombre_salida("Notas_Actividad1_%s.xlsx"))
    wb.save(salida)
    print("Escrito:", salida)
    print("Estudiantes:", len(nomina), " filas", primera, "a", ultima)
    print("Caras por cara en estadistica:", r - 5)
    return salida


if __name__ == "__main__":
    try:
        main()
    except contexto.FaltaArchivo as e:
        print("\nNo puedo generar la planilla todavia.\n")
        print(e)
        print("\nDeja la nomina de la seccion en la carpeta de arriba y vuelve a correr esto.")
        raise SystemExit(1)
