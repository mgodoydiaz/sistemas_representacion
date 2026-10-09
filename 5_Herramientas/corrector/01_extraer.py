#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
01_extraer.py -- Etapa E1 (extraccion) del corrector.

Lee el ZIP crudo descargado de Moodle, parsea los nombres de archivo segun el
patron que usa Moodle para las entregas de tarea, y extrae cada archivo a una
carpeta corta por alumno dentro de salida/entregas/<usuario>/, con nombres
numerados y truncados que evitan el error de Windows "ruta de acceso
demasiado larga" (0x80010135).

Tambien genera salida/inventario.csv cruzando la nomina oficial del curso con
lo que efectivamente se encontro en el ZIP.

Es idempotente: cada ejecucion recalcula el mapeo completo y deja las
carpetas de salida exactamente como corresponde, sin duplicar archivos.
"""

import csv
import io
import json
import os
import re
import shutil
import sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import contexto
import unicodedata
import zipfile
from pathlib import Path

# ---------------------------------------------------------------------------
# Rutas (segun CONTRATO.md: RAIZ = carpeta "Corrección Actividad 1")
# ---------------------------------------------------------------------------
RAIZ = Path(contexto.raiz())
# El zip y la nomina se detectan solos: asi este mismo script sirve en
# cualquier seccion sin editar rutas. Ver contexto.py.
ZIP_PATH = Path(contexto.ruta_zip())
NOMINA_PATH = Path(contexto.ruta_nomina())
SALIDA = RAIZ / "salida"
ENTREGAS = SALIDA / "entregas"
INVENTARIO_CSV = SALIDA / "inventario.csv"

MAX_SLUG_LEN = 40

EXT_PDF = {"pdf"}
EXT_DOC = {"docx", "doc"}
EXT_IMG = {"jpg", "jpeg", "png"}


# ---------------------------------------------------------------------------
# Utilidades de texto
# ---------------------------------------------------------------------------
def quitar_tildes(texto):
    """Elimina tildes y diacriticos, deja solo caracteres base."""
    nfkd = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in nfkd if not unicodedata.combining(c))


def slugify(nombre, maxlen=MAX_SLUG_LEN):
    """Convierte un nombre de archivo original en un slug corto, en
    minusculas, sin tildes ni caracteres raros, apto para nombre de archivo
    en cualquier sistema operativo."""
    if not nombre:
        return "archivo"
    s = quitar_tildes(nombre).lower()
    s = re.sub(r"[^a-z0-9]+", "_", s)
    s = re.sub(r"_+", "_", s).strip("_")
    if not s:
        s = "archivo"
    return s[:maxlen].strip("_") or "archivo"


def extraer_numero_lamina(nombre_original):
    """Busca el ultimo grupo de digitos en el nombre original (sin
    extension) para poder ordenar las entregas por numero de lamina, p.ej.
    'rec_vistas_color_3' -> 3. Devuelve None si no hay ningun digito."""
    if not nombre_original:
        return None
    grupos = re.findall(r"\d+", nombre_original)
    if not grupos:
        return None
    return int(grupos[-1])


def arreglar_mojibake(nombre, flag_utf8):
    """Moodle a veces genera ZIPs sin marcar el flag UTF-8 (bit 0x800) pero
    usando nombres UTF-8 igual; en ese caso Python los decodifica como
    cp437 y aparecen como mojibake. Si detectamos que el flag UTF-8 NO esta
    puesto, intentamos revertir la decodificacion cp437 y re-decodificar
    como utf-8; si el resultado es valido y "mejor" (sin caracteres de
    reemplazo raros), lo usamos."""
    if flag_utf8:
        return nombre
    try:
        crudo = nombre.encode("cp437")
        arreglado = crudo.decode("utf-8")
        return arreglado
    except (UnicodeDecodeError, UnicodeEncodeError):
        return nombre


# ---------------------------------------------------------------------------
# Parseo del patron de nombre de Moodle
# ---------------------------------------------------------------------------
# Patron general:
#   Actividad 1_<correo>_intento_<AAAA-MM-DD-HH-MM-SS>[_<nombre_original>].<ext>
# Los que NO traen "_<nombre_original>" y terminan en .txt son la ficha de
# metadatos del alumno que genera Moodle automaticamente.
PREFIJO = "Actividad %s_" % contexto.actividad()
MARCADOR_INTENTO = "_intento_"
RE_FECHA = re.compile(r"^\d{4}-\d{2}-\d{2}-\d{2}-\d{2}-\d{2}")


def parsear_nombre_moodle(nombre):
    """Devuelve un dict {correo, fecha, orig, ext} o None si el nombre no
    calza con el patron esperado de Moodle (archivo problematico/anomalo)."""
    if not nombre.startswith(PREFIJO):
        return None
    resto = nombre[len(PREFIJO):]
    idx = resto.find(MARCADOR_INTENTO)
    if idx == -1:
        return None
    correo = resto[:idx]
    resto2 = resto[idx + len(MARCADOR_INTENTO):]
    m = RE_FECHA.match(resto2)
    if not m:
        return None
    fecha = m.group(0)
    remanente = resto2[len(fecha):]
    if remanente.startswith("."):
        # Ficha de metadatos: "...<fecha>.txt" (sin nombre original)
        ext = remanente[1:]
        orig = None
    elif remanente.startswith("_"):
        origfull = remanente[1:]
        orig, ext = os.path.splitext(origfull)
        ext = ext[1:] if ext.startswith(".") else ext
    else:
        return None
    if "@" not in correo:
        return None
    return {"correo": correo.strip(), "fecha": fecha, "orig": orig, "ext": ext.lower()}


# ---------------------------------------------------------------------------
# Clasificacion del tipo de entrega
# ---------------------------------------------------------------------------
def clasificar_archivo(datos_bytes, nombre_original, ext):
    """Clasifica un archivo entregado segun su extension y, para imagenes,
    sus dimensiones (foto de celular vs. screenshot vs. captura del applet)."""
    ext_l = (ext or "").lower()
    if ext_l in EXT_PDF:
        return "pdf"
    if ext_l in EXT_DOC:
        return "docx"
    if ext_l in EXT_IMG:
        orig_l = (nombre_original or "").lower()
        if "rec_vistas_color" in orig_l:
            return "applet_jpg"
        try:
            from PIL import Image
            im = Image.open(io.BytesIO(datos_bytes))
            w, h = im.size
        except Exception:
            return "foto"
        if w and w > 2500:
            return "foto"
        if w and (h / w) > 1.5:
            return "screenshot"
        return "foto"
    return "otro"


def tipo_entrega_usuario(tipos_archivos):
    """A partir de la lista de tipos de cada archivo (sin la ficha) de un
    alumno, decide el tipo_entrega agregado para el CSV."""
    if not tipos_archivos:
        return "sin_archivos"
    distintos = set(tipos_archivos)
    if len(distintos) == 1:
        unico = distintos.pop()
        # El enunciado pide UN archivo PDF. Varios PDF sueltos se distinguen
        # aparte porque tienen su propio descuento.
        if unico == "pdf" and len(tipos_archivos) > 1:
            return "pdf_multiple"
        return unico
    return "mixto"


# ---------------------------------------------------------------------------
# Ficha de metadatos (.txt que genera Moodle)
# ---------------------------------------------------------------------------
RE_NOMBRE_FICHA = re.compile(r"^Nombre:\s*(.+?)\s*\(", re.MULTILINE)
RE_FECHA_FICHA = re.compile(r"^Fecha de env[ií]o:\s*(.+)$", re.MULTILINE)


def parsear_ficha(contenido):
    nombre = None
    fecha = None
    m = RE_NOMBRE_FICHA.search(contenido)
    if m:
        nombre = m.group(1).strip()
    m = RE_FECHA_FICHA.search(contenido)
    if m:
        fecha = m.group(1).strip()
    return nombre, fecha


# ---------------------------------------------------------------------------
# Lectura de la nomina oficial
# ---------------------------------------------------------------------------
def leer_nomina(path_nomina):
    """Lee la nomina (encabezados en la fila 5) y devuelve una lista de
    dicts ordenada segun aparece en la planilla."""
    import openpyxl

    wb = openpyxl.load_workbook(path_nomina, data_only=True)
    ws = contexto.hoja_nomina(wb)

    fila_encabezado = 5
    encabezados = {}
    for c in range(1, ws.max_column + 1):
        val = ws.cell(row=fila_encabezado, column=c).value
        if val:
            encabezados[str(val).strip()] = c

    col_n = encabezados.get("N°")
    col_apellidos = encabezados.get("Apellidos")
    col_nombres = encabezados.get("Nombres")
    col_nombre_completo = encabezados.get("Nombre completo")
    col_rut = encabezados.get("RUT")
    col_correo = encabezados.get("Correo institucional")

    alumnos = []
    for r in range(fila_encabezado + 1, ws.max_row + 1):
        correo = ws.cell(row=r, column=col_correo).value if col_correo else None
        if not correo:
            continue
        correo = str(correo).strip()
        nombre_completo = ws.cell(row=r, column=col_nombre_completo).value if col_nombre_completo else None
        alumnos.append({
            "n": ws.cell(row=r, column=col_n).value if col_n else None,
            "apellidos": ws.cell(row=r, column=col_apellidos).value if col_apellidos else None,
            "nombres": ws.cell(row=r, column=col_nombres).value if col_nombres else None,
            "nombre_completo": str(nombre_completo).strip() if nombre_completo else None,
            "rut": ws.cell(row=r, column=col_rut).value if col_rut else None,
            "correo": correo,
        })
    return alumnos


# ---------------------------------------------------------------------------
# Extraccion principal
# ---------------------------------------------------------------------------
def main():
    if not ZIP_PATH.exists():
        print(f"ERROR: no se encontro el ZIP en {ZIP_PATH}", file=sys.stderr)
        sys.exit(1)

    SALIDA.mkdir(parents=True, exist_ok=True)
    ENTREGAS.mkdir(parents=True, exist_ok=True)

    problematicos = []  # entradas del zip que no calzan con el patron esperado

    # ---- Paso 1: leer el zip y agrupar por correo ------------------------
    # entregas_por_correo[correo] = {
    #     "ficha": (ZipInfo, nombre_arreglado) | None,
    #     "archivos": [(ZipInfo, nombre_arreglado, orig, ext), ...],
    # }
    entregas_por_correo = {}

    with zipfile.ZipFile(ZIP_PATH, "r") as z:
        infolist = z.infolist()
        for info in infolist:
            flag_utf8 = bool(info.flag_bits & 0x800)
            nombre = arreglar_mojibake(info.filename, flag_utf8)
            datos = parsear_nombre_moodle(nombre)
            if datos is None:
                problematicos.append(nombre)
                continue
            correo = datos["correo"]
            entrada = entregas_por_correo.setdefault(
                correo, {"ficha": None, "fichas": [], "archivos": [],
                         "fecha_por_nombre": {}})
            entrada["fecha_por_nombre"][nombre] = datos["fecha"]
            if datos["orig"] is None and datos["ext"] == "txt":
                entrada["fichas"].append((info, nombre, datos["fecha"]))
                entrada["ficha"] = (info, nombre)
            else:
                entrada["archivos"].append((info, nombre, datos["orig"], datos["ext"]))

        # ---- Paso 2: por cada alumno, calcular el mapeo deseado ----------
        # mapeo_global[usuario] = {
        #     "correo": correo,
        #     "nombre_ficha": str|None, "fecha_ficha": str|None,
        #     "mapeo": {nombre_corto: nombre_original_en_zip},
        #     "archivos_info": [(nombre_corto, info, orig, ext)],
        # }
        mapeo_global = {}

        for correo, entrada in entregas_por_correo.items():
            usuario = correo.split("@")[0]
            carpeta_usuario = ENTREGAS / usuario

            # Moodle exporta TODOS los intentos del alumno, no solo el ultimo.
            # El que se corrige es el mas reciente: los anteriores suelen ser
            # subidas fallidas o parciales. Mezclarlos falsea el tipo de entrega
            # (por ejemplo, fotos de un primer intento + el PDF del segundo se
            # leerian como entrega "mixta"). Los previos se guardan aparte, en
            # intentos_previos/, para no perder la evidencia.
            fechas = set(entrada.get("fecha_por_nombre", {}).values())
            intentos_previos = []
            fichas_previas = []
            fecha_ultimo_intento = max(fechas) if fechas else None
            if len(fechas) > 1:
                intentos_previos = [
                    a for a in entrada["archivos"]
                    if entrada["fecha_por_nombre"].get(a[1]) != fecha_ultimo_intento]
                entrada["archivos"] = [
                    a for a in entrada["archivos"]
                    if entrada["fecha_por_nombre"].get(a[1]) == fecha_ultimo_intento]
                ultimas = [f for f in entrada.get("fichas", [])
                           if f[2] == fecha_ultimo_intento]
                if ultimas:
                    entrada["ficha"] = (ultimas[0][0], ultimas[0][1])
                # Las fichas de los intentos anteriores tambien se conservan:
                # traen el comentario que escribio el alumno al enviar, que a
                # veces explica por que reintento.
                fichas_previas = [f for f in entrada.get("fichas", [])
                                  if f[2] != fecha_ultimo_intento]

            # Ordenar archivos por numero de lamina si esta disponible, y
            # si no, alfabeticamente por el nombre original.
            archivos = entrada["archivos"]

            def clave_orden(item):
                _info, _nombre, orig, _ext = item
                num = extraer_numero_lamina(orig)
                clave_num = num if num is not None else 999999
                return (clave_num, (orig or "").lower())

            archivos_ordenados = sorted(archivos, key=clave_orden)

            mapeo = {}
            archivos_info = []
            for i, (info, nombre_zip, orig, ext) in enumerate(archivos_ordenados, start=1):
                slug = slugify(orig, MAX_SLUG_LEN)
                nombre_corto = f"{i:02d}_{slug}.{ext}" if ext else f"{i:02d}_{slug}"
                mapeo[nombre_corto] = nombre_zip
                archivos_info.append((nombre_corto, info, orig, ext))

            nombre_ficha_txt = None
            fecha_ficha_txt = None
            ficha_nombre_corto = None
            if entrada["ficha"] is not None:
                info_ficha, nombre_zip_ficha = entrada["ficha"]
                ficha_nombre_corto = "_ficha.txt"
                mapeo[ficha_nombre_corto] = nombre_zip_ficha
                contenido = z.read(info_ficha).decode("utf-8", errors="replace")
                nombre_ficha_txt, fecha_ficha_txt = parsear_ficha(contenido)

            mapeo_global[usuario] = {
                "correo": correo,
                "nombre_ficha": nombre_ficha_txt,
                "fecha_ficha": fecha_ficha_txt,
                "mapeo": mapeo,
                "archivos_info": archivos_info,
                "ficha_info": entrada["ficha"],
                "ficha_nombre_corto": ficha_nombre_corto,
                "intentos_previos": intentos_previos,
                "fichas_previas": fichas_previas,
                "n_intentos": len(fechas),
                "fecha_ultimo_intento": fecha_ultimo_intento,
            }

        # ---- Paso 3: escribir en disco (idempotente) ----------------------
        resumen_por_usuario = {}
        tipos_archivo_por_corto = {}  # para el CSV: usuario -> {nombre_corto: tipo}

        for usuario, info_usuario in mapeo_global.items():
            carpeta_usuario = ENTREGAS / usuario
            carpeta_usuario.mkdir(parents=True, exist_ok=True)

            deseados = set(info_usuario["mapeo"].keys())

            # Limpiar archivos que ya no correspondan (de una ejecucion
            # anterior con otra logica de nombres), para que el resultado
            # final sea siempre identico sin importar cuantas veces se
            # corra el script.
            for existente in list(carpeta_usuario.iterdir()):
                if existente.name == "mapeo.json":
                    continue
                if existente.name == "intentos_previos":
                    continue
                if existente.is_file() and existente.name not in deseados:
                    existente.unlink()

            tipos_por_corto = {}

            # Extraer archivos de contenido (no la ficha)
            for nombre_corto, info, orig, ext in info_usuario["archivos_info"]:
                datos_bytes = z.read(info)
                destino = carpeta_usuario / nombre_corto
                destino.write_bytes(datos_bytes)
                tipo = clasificar_archivo(datos_bytes, orig, ext)
                tipos_por_corto[nombre_corto] = tipo

            # Extraer la ficha
            if info_usuario["ficha_info"] is not None and info_usuario["ficha_nombre_corto"]:
                info_ficha, _nombre_zip_ficha = info_usuario["ficha_info"]
                destino_ficha = carpeta_usuario / info_usuario["ficha_nombre_corto"]
                destino_ficha.write_bytes(z.read(info_ficha))

            # Intentos anteriores: se guardan aparte, sin contarse como
            # entrega, y sin pasar por la clasificacion de tipo.
            if info_usuario.get("intentos_previos") or info_usuario.get("fichas_previas"):
                prev_dir = carpeta_usuario / "intentos_previos"
                prev_dir.mkdir(exist_ok=True)
                for info_f, nombre_zip_f, fecha_f in info_usuario.get("fichas_previas") or []:
                    destino_f = prev_dir / ("%s_ficha.txt" % fecha_f)
                    destino_f.write_bytes(z.read(info_f))
                    info_usuario["mapeo"]["intentos_previos/" + destino_f.name] = nombre_zip_f
                for j, (info_p, nombre_zip_p, orig_p, ext_p) in enumerate(
                        sorted(info_usuario["intentos_previos"], key=lambda t: t[1]), start=1):
                    fecha_p = nombre_zip_p.split("_intento_")[1][:19]
                    slug_p = slugify(orig_p, MAX_SLUG_LEN)
                    destino_p = prev_dir / ("%s_%02d_%s.%s" % (fecha_p, j, slug_p, ext_p))
                    destino_p.write_bytes(z.read(info_p))
                    info_usuario["mapeo"]["intentos_previos/" + destino_p.name] = nombre_zip_p

            # Escribir mapeo.json (trazabilidad nombre_corto -> nombre en el zip)
            mapeo_json_path = carpeta_usuario / "mapeo.json"
            mapeo_json_path.write_text(
                json.dumps(info_usuario["mapeo"], ensure_ascii=False, indent=2, sort_keys=True),
                encoding="utf-8",
            )

            tipos_archivo_por_corto[usuario] = tipos_por_corto
            resumen_por_usuario[usuario] = {
                "correo": info_usuario["correo"],
                "n_archivos": len(info_usuario["archivos_info"]),
                "archivos": [nc for nc, _i, _o, _e in info_usuario["archivos_info"]],
                "nombre_ficha": info_usuario["nombre_ficha"],
                "fecha_ficha": info_usuario["fecha_ficha"],
                "tipos": list(tipos_por_corto.values()),
                "n_intentos": info_usuario.get("n_intentos", 1),
                "n_previos": len(info_usuario.get("intentos_previos") or []),
            }

    # ---- Paso 4: cruzar con la nomina y construir inventario.csv ---------
    alumnos_nomina = leer_nomina(NOMINA_PATH)

    filas_csv = []
    usuarios_nomina = set()
    for alumno in alumnos_nomina:
        correo = alumno["correo"]
        usuario = correo.split("@")[0]
        usuarios_nomina.add(usuario)
        datos_zip = resumen_por_usuario.get(usuario)

        observaciones = []

        if datos_zip is None:
            filas_csv.append({
                "usuario": usuario,
                "correo": correo,
                "nombre_completo": alumno["nombre_completo"] or "",
                "fecha_envio": "",
                "n_archivos": 0,
                "archivos": "",
                "tipo_entrega": "sin_archivos",
                "observacion": "No entrego archivos (no aparece en el ZIP de Moodle)",
            })
            continue

        n_archivos = datos_zip["n_archivos"]
        tipo = tipo_entrega_usuario(datos_zip["tipos"])
        nombre_completo = datos_zip["nombre_ficha"] or alumno["nombre_completo"] or ""

        if n_archivos == 0:
            observaciones.append("Entrego la ficha pero sin archivos adjuntos")
        if tipo == "mixto":
            observaciones.append(f"Mezcla de tipos de archivo: {sorted(set(datos_zip['tipos']))}")
        if "otro" in datos_zip["tipos"]:
            observaciones.append("Incluye extension de archivo no reconocida")
        if tipo == "docx":
            observaciones.append("Entrego un documento Word (.docx) en vez de imagen o PDF")
        if datos_zip.get("n_intentos", 1) > 1:
            observaciones.append(
                "Hizo %d intentos; se corrige el ultimo. Los %d archivo(s) de "
                "intentos anteriores quedaron en intentos_previos/"
                % (datos_zip["n_intentos"], datos_zip.get("n_previos", 0)))

        filas_csv.append({
            "usuario": usuario,
            "correo": correo,
            "nombre_completo": nombre_completo,
            "fecha_envio": datos_zip["fecha_ficha"] or "",
            "n_archivos": n_archivos,
            "archivos": ";".join(datos_zip["archivos"]),
            "tipo_entrega": tipo,
            "observacion": " | ".join(observaciones),
        })

    # Alumnos que entregaron pero no estan en la nomina oficial (anomalia)
    extra_no_nomina = []
    for usuario, datos_zip in resumen_por_usuario.items():
        if usuario in usuarios_nomina:
            continue
        tipo = tipo_entrega_usuario(datos_zip["tipos"])
        filas_csv.append({
            "usuario": usuario,
            "correo": datos_zip["correo"],
            "nombre_completo": datos_zip["nombre_ficha"] or "",
            "fecha_envio": datos_zip["fecha_ficha"] or "",
            "n_archivos": datos_zip["n_archivos"],
            "archivos": ";".join(datos_zip["archivos"]),
            "tipo_entrega": tipo,
            "observacion": "Correo no encontrado en la nomina oficial del curso",
        })
        extra_no_nomina.append(usuario)

    with open(INVENTARIO_CSV, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=[
            "usuario", "correo", "nombre_completo", "fecha_envio",
            "n_archivos", "archivos", "tipo_entrega", "observacion",
        ])
        writer.writeheader()
        for fila in filas_csv:
            writer.writerow(fila)

    # ---- Paso 5: resumen por consola --------------------------------------
    total_extraido = sum(d["n_archivos"] for d in resumen_por_usuario.values())
    print("=" * 70)
    print("RESUMEN DE EXTRACCION - Actividad %s %s" % (contexto.actividad(), contexto.codigo_seccion()))
    print("=" * 70)
    print(f"Alumnos en la nomina oficial: {len(alumnos_nomina)}")
    print(f"Alumnos que entregaron (encontrados en el ZIP): {len(resumen_por_usuario)}")
    print(f"Alumnos sin entrega: {len(alumnos_nomina) - (len(resumen_por_usuario) - len(extra_no_nomina))}")
    print(f"Total de archivos de contenido extraidos: {total_extraido}")
    print()
    print("Detalle por usuario:")
    for usuario in sorted(resumen_por_usuario.keys()):
        d = resumen_por_usuario[usuario]
        tipo = tipo_entrega_usuario(d["tipos"])
        print(f"  - {usuario:35s} {d['n_archivos']} archivo(s)  tipo={tipo}")
    if extra_no_nomina:
        print()
        print("ALERTA: correos que entregaron pero NO estan en la nomina:")
        for u in extra_no_nomina:
            print(f"  - {u}")
    if problematicos:
        print()
        print("ALERTA: entradas del ZIP que no calzaron con el patron esperado de Moodle:")
        for p in problematicos:
            print(f"  - {p}")
    print()
    print(f"inventario.csv escrito en: {INVENTARIO_CSV}")
    print(f"Carpetas de entregas en:   {ENTREGAS}")
    print("=" * 70)


if __name__ == "__main__":
    main()
