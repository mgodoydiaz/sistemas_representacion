"""Genera los enunciados y pautas de las Actividades 6 (Cortes) y 7 (Isometrico).
Uso:  python3 figuras_actividades.py && python3 build_actividades.py
Requiere la skill plantilla-miguel (scripts/plantilla.py) y LibreOffice."""
import os
import sys
import re

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL = os.environ.get("PLANTILLA_MIGUEL", "/mnt/skills/plugins/plantilla-miguel")
sys.path.insert(0, os.path.join(SKILL, "scripts"))
from plantilla import Doc
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

FIG = os.path.join(HERE, "figuras")
F = lambda n: os.path.join(FIG, n + ".png")


def rich(d, text, style="Normal", before=None, after=None, align=None):
    """Parrafo con **negrita** en linea."""
    p = d.doc.add_paragraph(style=style)
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            p.add_run(part).bold = (i % 2 == 1)
    if before is not None:
        p.paragraph_format.space_before = Pt(before)
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    return p


def bullets(d, items):
    for it in items:
        p = rich(d, "•\t" + it, after=1)
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.first_line_indent = Inches(-0.2)


def images(d, paths_widths, caption=None):
    """Una o varias figuras en una fila. Si son varias, se componen en un solo PNG
    (figuras/_fila_*.png) para que queden lado a lado en Word y en LibreOffice."""
    if len(paths_widths) > 1:
        from PIL import Image
        ims = [Image.open(p).convert("RGB") for p, _ in paths_widths]
        dpi = 300
        ims = [im.resize((int(w * dpi), int(w * dpi * im.height / im.width)), Image.LANCZOS)
               for im, (_, w) in zip(ims, paths_widths)]
        gap = int(0.15 * dpi)
        W = sum(im.width for im in ims) + gap * (len(ims) - 1)
        Hh = max(im.height for im in ims)
        canvas = Image.new("RGB", (W, Hh), "white")
        x = 0
        for im in ims:
            canvas.paste(im, (x, (Hh - im.height) // 2))
            x += im.width + gap
        name = "_fila_" + "_".join(os.path.basename(p)[:-4] for p, _ in paths_widths) + ".png"
        path = os.path.join(FIG, name)
        canvas.save(path, dpi=(dpi, dpi))
        paths_widths = [(path, W / dpi)]
    p = d.doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for path, w in paths_widths:
        p.add_run().add_picture(path, width=Inches(w))
    if caption:
        c = d.doc.add_paragraph(caption, style="Caption")
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER


def table(d, headers, rows, widths=None, small=True):
    d.table(headers, rows)
    t = d.doc.tables[-1]
    t.autofit = False
    if widths:
        for i, col in enumerate(t.columns):
            col.width = Inches(widths[i])
    for r in t.rows:
        for i, c in enumerate(r.cells):
            if widths:
                c.width = Inches(widths[i])
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                if i >= 2:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.size = Pt(9.5 if small else 10.5)
    d.doc.add_paragraph().paragraph_format.space_after = Pt(0)


def new(rhead):
    d = Doc()
    d.headers(lhead="Sistemas de Representación · PCI 1119", rhead=rhead,
              lfoot="Ingeniería Civil Industrial · UCT", cfoot="Prof. Miguel Godoy Díaz")
    return d


SUB = "Actividad en clase · individual · 80 minutos · Fecha: por definir"


# ============================================================ ACTIVIDAD 6
def actividad_6():
    d = new("Actividad 6 · Cortes")
    d.title("Actividad 6: Cortes", SUB)
    rich(d, "**Objetivo.** Representar el interior de una pieza mediante corte total y "
            "semicorte, con la indicación y el rayado que pide la norma ISO 128.")
    rich(d, "**Materiales.** Hojas blancas tamaño carta o A4, portaminas de dos grosores, "
            "regla, escuadras y compás.")
    rich(d, "**Instrucciones.** Trabaje a mano y en escala 1:1. Use una hoja por ejercicio, "
            "con margen y cajetín (título, escala, autor, fecha). Las medidas están en "
            "milímetros. Las vistas dadas siguen el primer diedro (ISO-E). No es necesario "
            "acotar sus dibujos. Use la notación del apunte, capítulo 5.")
    d.h1("Ejercicio 1: buje con brida (30 puntos)")
    rich(d, "La Figura 1 muestra el alzado y la planta de un buje con brida. Dibuje la planta "
            "y, sobre ella, el alzado en **corte total A-A** por el eje de la pieza. Indique "
            "la traza del plano de corte en la planta.")
    d.h1("Ejercicio 2: soporte con nervios (35 puntos)")
    rich(d, "La Figura 2 muestra un soporte simétrico: base con dos agujeros, cilindro "
            "central con agujero pasante y dos nervios de refuerzo de espesor 6. Dibuje la "
            "planta y el alzado en **semicorte**: mitad izquierda en vista y mitad derecha "
            "cortada. Indique la traza en la planta. Recuerde cómo se representa un nervio "
            "cortado a lo largo.")
    images(d, [(F("A6_ej1_buje"), 2.6), (F("A6_ej2_soporte_nervio"), 3.7)],
           "Figura 1. Buje con brida (izquierda). Figura 2. Soporte con nervios (derecha).")
    d.h1("Ejercicio 3: cuerpo de válvula simplificado (35 puntos)")
    rich(d, "La Figura 3 muestra el cuerpo de una válvula de paso recto, formado solo por "
            "cilindros: dos bridas de conexión, dos cuellos, un cuerpo central y una boca "
            "superior con brida para el bonete. Por dentro tiene un conducto Ø20 pasante, "
            "una cámara central Ø30 de largo 40 y un agujero vertical Ø30 que llega desde la "
            "boca superior hasta el eje del conducto. Dibuje la vista lateral y el alzado en "
            "**corte total A-A** por el plano longitudinal de simetría. Indique la traza en "
            "la vista lateral.")
    images(d, [(F("A6_ej3_valvula"), 6.4)],
           "Figura 3. Cuerpo de válvula simplificado: alzado y vista lateral izquierda.")
    rich(d, "Simplificaciones: la pieza no tiene agujeros para pernos, redondeos ni asiento "
            "interior. Los cilindros de igual diámetro que se cruzan (Ø50 por fuera y Ø30 "
            "por dentro) se cortan en dos curvas planas, que en el alzado se ven como rectas "
            "a 45°.")
    d.h1("Qué se evalúa")
    bullets(d, [
        "Elección y posición del plano de corte.",
        "Indicación del corte: traza, tramos gruesos, flechas en la dirección de observación y letras.",
        "Rótulo A-A sobre la vista cortada.",
        "Rayado fino a 45°, uniforme, solo en el material cortado y hasta el contorno.",
        "Vista cortada sin aristas ocultas.",
        "Elementos que no se rayan (nervios y ejes) y semicorte separado por línea de eje.",
        "Correspondencia entre vistas, limpieza y grosores de línea.",
    ])
    d.h1("Entrega")
    rich(d, "Entregue sus hojas con cajetín al final de la clase. La nota se promedia en el "
            "30% de Actividades.")
    d.save(os.path.join(HERE, "Actividad_6_Cortes.docx"), pdf=True)


def pauta_6():
    d = new("Pauta Actividad 6 · Cortes")
    d.title("Pauta Actividad 6: Cortes", "Solución y puntaje · uso del docente")
    rich(d, "Puntaje total: 100 puntos (Ejercicio 1: 30, Ejercicio 2: 35, Ejercicio 3: 35). "
            "Las anotaciones en azul son comentarios de corrección y no forman parte del "
            "dibujo que entrega el estudiante.")
    d.h1("Ejercicio 1: buje con brida, corte total A-A")
    images(d, [(F("A6_ej1_buje_solucion"), 3.7)], "Figura 1. Solución del ejercicio 1.")
    bullets(d, [
        "El plano pasa por el eje y por dos de los cuatro agujeros Ø10, que aparecen cortados.",
        "Se rayan cuatro zonas: dos del cubo con la brida y dos franjas exteriores de la brida.",
        "No hay línea entre el cubo y la brida dentro del rayado: es una sola pieza.",
        "Se acepta la traza por el eje vertical de la planta si la vista cortada es coherente con ella.",
    ])
    d.h1("Ejercicio 2: soporte con nervios, semicorte")
    images(d, [(F("A6_ej2_soporte_nervio_solucion"), 4.6)], "Figura 2. Solución del ejercicio 2.")
    bullets(d, [
        "Mitad izquierda en vista exterior, sin ocultas. Mitad derecha cortada por el eje.",
        "El nervio queda sin rayar, limitado por su contorno. La base y el cilindro sí se rayan.",
        "El límite entre mitades es la línea de eje. Si se dibuja gruesa, se pierde C7.",
        "El semicorte no lleva letras ni rótulo: traza en L con una flecha (apunte, Figura 12). "
        "Si el estudiante agrega letras y rótulo coherentes, no se descuenta.",
        "Si el estudiante hace corte total con ambos nervios sin rayar, asigne C1 a C6 con normalidad y 0 en C7.",
    ])
    d.h1("Ejercicio 3: cuerpo de válvula, corte total A-A")
    images(d, [(F("A6_ej3_valvula_solucion"), 6.0)], "Figura 3. Solución del ejercicio 3.")
    bullets(d, [
        "La zona rayada es una sola (pieza de una sola parte), con el mismo ángulo y separación.",
        "Quedan en blanco el conducto Ø20, la cámara Ø30 y el agujero vertical Ø30.",
        "Detrás del plano se ven las rectas a 45° donde se cruzan los agujeros Ø30, y los bordes "
        "del conducto en las caras de la cámara y de las bridas.",
        "La V exterior del cruce de los cilindros Ø50 desaparece: queda delante del plano.",
        "La traza va en la vista lateral, vertical por el eje, con las flechas hacia la izquierda "
        "(hacia el fondo de la pieza).",
    ])
    d.h1("Tabla de puntaje")
    rows = [
        ["C1", "Elección y posición del plano de corte", 3, 3, 4, 10],
        ["C2", "Indicación: traza, tramos gruesos, flechas en la dirección de observación, letras", 6, 5, 6, 17],
        ["C3", "Rótulo A-A sobre la vista cortada", 2, 0, 3, 5],
        ["C4", "Rayado a 45°, fino, uniforme, solo material cortado, termina en el contorno", 8, 7, 9, 24],
        ["C5", "Sin aristas ocultas en la vista cortada", 5, 5, 6, 16],
        ["C6", "Elementos que no se rayan (nervios, ejes)", 0, 6, 0, 6],
        ["C7", "Semicorte separado por línea de eje", 0, 5, 0, 5],
        ["C8", "Correspondencia entre vistas, limpieza y grosores", 6, 4, 7, 17],
        ["", "Total", 30, 35, 35, 100],
    ]
    table(d, ["Crit.", "Qué se exige", "Ej. 1", "Ej. 2", "Ej. 3", "Total"], rows,
          widths=[0.55, 3.75, 0.55, 0.55, 0.55, 0.55])
    d.h1("Descuentos frecuentes")
    table(d, ["Error", "Criterio", "Descuento sugerido"], [
        ["Flechas en sentido contrario a la vista dibujada", "C2", "mitad del criterio"],
        ["Sin tramos gruesos en los extremos de la traza", "C2", "2 puntos"],
        ["Rayado dentro de agujeros o huecos", "C4", "mitad del criterio"],
        ["Rayado con dirección o separación distinta en zonas de la misma pieza", "C4", "3 puntos"],
        ["Aristas ocultas en la vista cortada", "C5", "2 puntos por zona, hasta el total"],
        ["Faltan aristas visibles detrás del plano", "C8", "2 puntos por arista"],
        ["Nervio rayado", "C6", "todo el criterio"],
    ], widths=[3.6, 0.9, 2.0])
    d.save(os.path.join(HERE, "Pauta_Actividad_6_Cortes.docx"), pdf=True)


# ============================================================ ACTIVIDAD 7
def actividad_7():
    d = new("Actividad 7 · Isométrico")
    d.title("Actividad 7: Isométrico", SUB)
    rich(d, "**Objetivo.** Construir el dibujo isométrico de una pieza a partir de sus tres "
            "vistas, incluyendo planos inclinados, círculos y arcos.")
    rich(d, "**Materiales.** Plantilla isométrica del curso (de 5 mm o de 10 mm), portaminas "
            "de dos grosores, regla y escuadras. Puede usar plantilla de elipses o compás.")
    rich(d, "**Instrucciones.** Trabaje a mano y en escala 1:1 sobre la plantilla isométrica. "
            "Las medidas están en milímetros y las vistas siguen el primer diedro (ISO-E). "
            "Dibuje cada pieza vista desde el frente, la izquierda y arriba, de modo que se "
            "vean las caras del alzado, de la planta y de la lateral izquierda. Dibuje solo "
            "las aristas visibles. No es necesario acotar. Agregue un cajetín con título, "
            "escala, autor y fecha.")
    d.h1("Pieza 1: pieza escalonada (25 puntos)")
    rich(d, "Dibuje el isométrico de la pieza de la Figura 1. Todas sus caras son planas y "
            "paralelas a los planos principales.")
    d.h1("Pieza 2: pieza con plano inclinado (30 puntos)")
    rich(d, "Dibuje el isométrico de la pieza de la Figura 2. Tiene un plano inclinado y una "
            "ranura superior de 20 de ancho cuyo fondo está a 30 de altura. Recuerde que un "
            "plano inclinado se resuelve ubicando sus extremos y uniéndolos.")
    images(d, [(F("A7_p1_escalonada"), 3.05), (F("A7_p2_plano_inclinado"), 3.05)],
           "Figura 1. Pieza escalonada (izquierda). Figura 2. Pieza con plano inclinado (derecha).")
    d.h1("Pieza 3: soporte con respaldo redondeado (45 puntos)")
    rich(d, "Dibuje el isométrico del soporte de la Figura 3. La base mide 40 × 60 × 10 y "
            "tiene un agujero vertical pasante Ø10. El respaldo, de espesor 10, termina en "
            "un semicírculo R20 con un agujero pasante Ø20 concéntrico. Dibuje los círculos "
            "como elipses en el isoplano que corresponde y cuide las tangencias.")
    images(d, [(F("A7_p3_soporte_respaldo"), 4.1)],
           "Figura 3. Soporte con respaldo redondeado.")
    d.h1("Qué se evalúa")
    bullets(d, [
        "Ejes isométricos a 30° y verticales.",
        "Medidas tomadas solo sobre las direcciones isométricas.",
        "Planos inclinados resueltos uniendo sus extremos.",
        "Círculos dibujados como elipses en el isoplano correcto.",
        "Tangencias y arcos bien resueltos.",
        "Solo aristas visibles.",
        "Proporción, limpieza y grosores de línea.",
    ])
    d.h1("Entrega")
    rich(d, "Entregue sus hojas con cajetín al final de la clase. La nota se promedia en el "
            "30% de Actividades. Guarde una copia de la pieza 3: se volverá a usar en AutoCAD.")
    d.save(os.path.join(HERE, "Actividad_7_Isometrico.docx"), pdf=True)


def pauta_7():
    d = new("Pauta Actividad 7 · Isométrico")
    d.title("Pauta Actividad 7: Isométrico", "Solución y puntaje · uso del docente")
    rich(d, "Puntaje total: 100 puntos (Pieza 1: 25, Pieza 2: 30, Pieza 3: 45). Las soluciones "
            "están dibujadas sobre una retícula isométrica de 10 mm, vistas desde el frente, la "
            "izquierda y arriba. Las anotaciones en azul son comentarios de corrección. Acepte "
            "también la vista desde la derecha si el dibujo es coherente con las vistas dadas.")
    d.h1("Piezas 1 y 2")
    images(d, [(F("A7_p1_escalonada_solucion"), 3.05), (F("A7_p2_plano_inclinado_solucion"), 3.05)],
           "Figura 1. Solución de la pieza 1 (izquierda) y de la pieza 2 (derecha).")
    bullets(d, [
        "Pieza 1: tres niveles a 10, 25 y 40 de altura. El bloque superior ocupa solo la mitad del fondo.",
        "Pieza 2: el plano inclinado va de (X = 5, Z = 15) a (X = 35, Z = 40). Sus aristas no son "
        "isométricas, por lo que su largo no se mide: resultan de unir los extremos.",
        "Pieza 2: el fondo de la ranura (Z = 30) corta al plano inclinado en X = 23. Ese punto se "
        "obtiene en el dibujo, no se mide.",
    ])
    d.h1("Pieza 3: soporte con respaldo redondeado")
    images(d, [(F("A7_p3_soporte_respaldo_solucion"), 5.0)], "Figura 2. Solución de la pieza 3.")
    bullets(d, [
        "Agujero Ø10: elipse en el isoplano horizontal, centrada a 20 del costado y a 40 del fondo. "
        "No se ve su borde inferior.",
        "Respaldo: dos semielipses R20 (cara delantera y cara trasera) unidas por una recta tangente "
        "en la dirección de la profundidad. La cara trasera solo se ve en parte.",
        "Agujero Ø20: elipse completa en la cara delantera. A través del agujero se ve un arco del "
        "borde trasero.",
        "Los arcos nacen a la altura 30, tangentes a las aristas verticales del respaldo.",
    ])
    d.h1("Tabla de puntaje")
    rows = [
        ["I1", "Ejes isométricos a 30° y verticales", 6, 5, 5, 16],
        ["I2", "Medidas tomadas solo sobre direcciones isométricas", 8, 7, 6, 21],
        ["I3", "Planos inclinados resueltos uniendo extremos", 0, 8, 0, 8],
        ["I4", "Círculos como elipses en el isoplano correcto", 0, 0, 14, 14],
        ["I5", "Tangencias y arcos", 0, 0, 8, 8],
        ["I6", "Solo aristas visibles", 6, 5, 6, 17],
        ["I7", "Proporción, limpieza y grosores", 5, 5, 6, 16],
        ["", "Total", 25, 30, 45, 100],
    ]
    table(d, ["Crit.", "Qué se exige", "Pieza 1", "Pieza 2", "Pieza 3", "Total"], rows,
          widths=[0.55, 3.55, 0.65, 0.65, 0.65, 0.55])
    d.doc.add_page_break()
    d.h1("Descuentos frecuentes")
    table(d, ["Error", "Criterio", "Descuento sugerido"], [
        ["Medir el largo real sobre la arista inclinada", "I2, I3", "mitad de I3"],
        ["Círculo dibujado como circunferencia", "I4", "todo el criterio en ese agujero"],
        ["Elipse en el isoplano equivocado (eje mayor mal orientado)", "I4", "5 puntos por elipse"],
        ["Falta la recta tangente entre los dos arcos del respaldo", "I5", "4 puntos"],
        ["Arco que no llega tangente a las aristas verticales", "I5", "2 puntos"],
        ["Aristas ocultas dibujadas", "I6", "2 puntos por zona"],
        ["Falta el arco trasero visible a través del agujero Ø20", "I6", "2 puntos"],
    ], widths=[3.6, 0.9, 2.0])
    d.save(os.path.join(HERE, "Pauta_Actividad_7_Isometrico.docx"), pdf=True)


if __name__ == "__main__":
    for f in (actividad_6, pauta_6, actividad_7, pauta_7):
        f()
        print("ok", f.__name__)
