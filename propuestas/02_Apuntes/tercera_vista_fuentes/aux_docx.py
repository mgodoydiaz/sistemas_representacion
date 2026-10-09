from docx import Document
import subprocess
src='/home/claude/sistemas_representacion/2_Clases/Apuntes/Nuevo_Cajetin_Vistas_Cortes/Apuntes_Unidad1_PCI1119_v3.docx'
d=Document(src)
ps=d.paragraphs
ps[0].runs[0].text='Vistas auxiliares, cortes y secciones'
for r in ps[0].runs[1:]: r.text=''
start=ps[157]._p; body=d.element.body
keep={ps[0]._p,ps[1]._p}
for el in list(body):
    if el is start: break
    if el in keep: continue
    body.remove(el)
# encabezado derecho
for s in d.sections:
    for part in (s.header,s.footer):
        for p in part.paragraphs:
            for r in p.runs:
                if 'Unidad 1' in r.text: r.text=r.text.replace('Unidad 1 · Fundamentos de la representación','Vistas auxiliares y cortes')
d.save('Apunte_Vistas_Auxiliares_y_Cortes.docx')
subprocess.run(['soffice','--headless','--convert-to','pdf','Apunte_Vistas_Auxiliares_y_Cortes.docx'],capture_output=True)
