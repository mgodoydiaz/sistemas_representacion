"""Arma el apunte "Dibujo isométrico en AutoCAD" (docx + pdf) con la plantilla de Miguel.
Uso: python3 figuras_isocad.py && python3 build_isocad.py
"""
import os
import sys

sys.path.insert(0, "/mnt/skills/plugins/plantilla-miguel/scripts")
from plantilla import Doc
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
FIG = os.path.join(HERE, "figuras")
OUT = os.path.join(HERE, "Apunte_Isometrico_en_AutoCAD.docx")

d = Doc()


def tabla(headers, rows, widths, size=9.5):
    """Tabla del molde con letra compacta, celdas a la izquierda y anchos fijos (pulgadas)."""
    if d.doc.paragraphs and d.doc.paragraphs[-1].style.name != "Caption":
        d.doc.paragraphs[-1].paragraph_format.keep_with_next = True   # la frase previa viaja con la tabla
    d.table(headers, rows)
    t = d.doc.tables[-1]
    t.autofit = False
    for i, col in enumerate(t.columns):
        col.width = Inches(widths[i])
    last = len(t.rows) - 1
    for k, row in enumerate(t.rows):
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))        # la fila no se parte
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i])
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.keep_with_next = k < last   # la tabla no se parte
                p.paragraph_format.space_before = Pt(1)
                p.paragraph_format.space_after = Pt(1)
                for r in p.runs:
                    r.font.size = Pt(size)
    sp = d.doc.add_paragraph()                      # separación mínima tras la tabla
    sp.paragraph_format.space_after = Pt(0)
    sp.paragraph_format.line_spacing = Pt(5)
    return t


def fig(name, width, caption):
    d.image(os.path.join(FIG, name), width_in=width, caption=caption)
    d.doc.paragraphs[-2].paragraph_format.keep_with_next = True
    d.doc.paragraphs[-2].paragraph_format.space_after = Pt(2)


d.title("Dibujo isométrico en AutoCAD",
        "Apunte de clase · PCI 1119 Sistemas de Representación")

# ------------------------------------------------------------------ 1
d.h1("El problema")
d.body("Usted ya dibuja isométricos a mano: ejes a 30°, método de la caja y círculos "
       "como elipses. En AutoCAD 2D la hoja es plana y el programa no sabe que su dibujo "
       "es isométrico. CIRCLE traza siempre un círculo perfecto y las cotas salen "
       "perpendiculares a la arista, aunque la cara esté inclinada.")
d.destacado(2, "Dibujo isométrico 2D: líneas y elipses planas que imitan el volumen. "
               "Modelo 3D: un sólido que se puede girar. En esta clase se trabaja solo el primero.")
d.body("La solución es el modo de dibujo isométrico (ISODRAFT). Los comandos de este apunte "
       "van en inglés. En AutoCAD en español escríbalos con un guion bajo delante, por "
       "ejemplo _ISODRAFT o _ELLIPSE.")

# ------------------------------------------------------------------ 2
d.h1("Preparación del archivo")
d.body("Con UNITS deje el tipo Decimal y las unidades en milímetros. Dibuje siempre 1:1: "
       "40 mm son 40 unidades. La escala se decide al imprimir. Cree las capas con LAYER, "
       "con los grosores de ISO 128 que ya conoce:")
tabla(["Capa", "Qué contiene", "Tipo de línea", "Grosor"],
      [["Visible", "Aristas y contornos vistos", "Continuous", "0,50 mm"],
       ["Ejes", "Ejes de agujeros y arcos", "CENTER (trazo y punto)", "0,25 mm"],
       ["Cotas", "Cotas y textos", "Continuous", "0,25 mm"],
       ["Auxiliar", "Líneas de construcción; se borra o se apaga al final", "Continuous", "0,25 mm"]],
      [1.0, 2.9, 1.6, 1.0])
d.body("Active la referencia a objetos (F3) con Endpoint, Midpoint, Center, Quadrant e "
       "Intersection. Sin ella los trazos no cierran.")

# ------------------------------------------------------------------ 3
d.h1("Activar el modo isométrico")
d.body("Escriba ISODRAFT y elija un isoplano: Left, Top o Right. La opción Orthographic lo "
       "apaga. También puede usar el botón Isometric Drafting de la barra de estado. En "
       "versiones antiguas se usa la variable SNAPSTYL: 1 lo activa y 0 lo desactiva.")
d.destacado(2, "Isoplano: cada una de las tres caras del cubo isométrico. El isoplano activo "
               "define sobre qué dos ejes se mueve el cursor y cómo se orientan los isocírculos.")
d.body("Con F5 o Ctrl+E se pasa al isoplano siguiente. Con Ortho (F8) el cursor queda "
       "amarrado a los dos ejes del isoplano activo. Otra opción es el rastreo polar (F10) "
       "con incremento de 30°.")
tabla(["Isoplano", "Ejes disponibles", "Caras que se dibujan", "En la pieza guía"],
      [["Left (izquierdo)", "90° y 150°", "Verticales que miran a la izquierda",
        "Frente de la base y cara delantera del respaldo (R20 y Ø20)"],
       ["Top (superior)", "30° y 150°", "Horizontales", "Cara superior de la base (Ø10)"],
       ["Right (derecho)", "30° y 90°", "Verticales que miran a la derecha",
        "Costado derecho de la base y del respaldo"]],
      [1.25, 1.15, 1.9, 2.2])
fig("isocad_01_isoplanos.png", 6.4,
    "Figura 1. Ejes isométricos, isoplanos de AutoCAD y trazado de una arista inclinada.")

# ------------------------------------------------------------------ 4
d.h1("Líneas")
d.body("Use LINE con entrada directa de distancia: haga clic en el punto inicial, mueva "
       "el cursor en la dirección deseada, escriba la medida y pulse Enter. Con Ortho "
       "activo la dirección queda fija sobre un eje isométrico.")
d.destacado(2, "Solo las aristas paralelas a los ejes (30°, 90°, 150°) conservan su medida "
               "real. Una arista inclinada se ve deformada: no se mide, se construye.")
d.body("Para una arista inclinada, ubique sus dos extremos midiendo sobre los ejes. Luego "
       "apague Ortho (F8) y únalos con LINE y Endpoint (figura 1c).")

# ------------------------------------------------------------------ 5
d.h1("Círculos, arcos y tangencias")
d.body("Un círculo sobre una cara isométrica se ve como elipse: es el isocírculo. Elija "
       "primero el isoplano de la cara (F5). Luego escriba ELLIPSE, opción Isocircle (I), "
       "marque el centro y dé el radio (o D y el diámetro). La opción Isocircle solo "
       "aparece con el modo isométrico activo.")
fig("isocad_02_circle_vs_isocircle.png", 3.6,
    "Figura 2. CIRCLE no sirve en un isométrico; cada cara necesita su isocírculo.")
tabla(["Caso", "Cómo se resuelve"],
      [["Arco o redondeo", "Dibuje el isocírculo completo y recorte lo que sobra con TRIM. "
                           "FILLET no sirve: genera arcos circulares."],
       ["Círculos concéntricos", "Otro isocírculo con el mismo centro. No use OFFSET sobre una elipse."],
       ["Tangencia con aristas", "Un isocírculo de radio R cuyo centro está a R de cada arista "
                                 "queda tangente a ellas."],
       ["Cara trasera", "COPY del isocírculo, con un desplazamiento igual al espesor sobre "
                        "el eje de la profundidad."],
       ["Línea tangente entre dos isocírculos", "LINE entre los puntos Quadrant de los extremos "
                                                "del eje mayor de cada elipse."]],
      [1.9, 4.6])

# ------------------------------------------------------------------ 6
d.h1("Ejemplo resuelto: soporte con respaldo redondeado")
d.body("Es la pieza que usted dibujó a mano en la actividad anterior. La base mide "
       "40 × 60 × 10 y el respaldo 40 × 10 × 50, con el extremo superior redondeado (R20).")
fig("isocad_03_pieza_guia.png", 3.9,
    "Figura 3. Pieza guía en isométrico. X es el ancho, Y la profundidad y Z la altura.")
tabla(["Paso", "Comando", "Qué se hace y qué se obtiene"],
      [["1", "ISODRAFT, LINE, F5",
        "Caja de la base desde el vértice A: 40 hacia 150°, 60 hacia 30° y 10 hacia 90°. "
        "Cambie de isoplano para cada cara."],
       ["2", "LINE",
        "Caja del respaldo al fondo: 10 de espesor y 40 de alto sobre la base."],
       ["3", "F5 (Left), LINE, ELLIPSE (I)",
        "Centro C: desde el punto medio (Midpoint) del pie del respaldo, suba 20. "
        "Con centro C, isocírculos de radio 20 y de radio 10."],
       ["4", "COPY, LINE",
        "Copie ambos isocírculos 10 mm hacia 30° (cara trasera). Una con LINE los "
        "Quadrant superiores izquierdos de los dos R20: es la línea tangente."],
       ["5", "TRIM, ERASE",
        "Borre la mitad inferior de los R20, las esquinas de la caja sobre el arco, "
        "las líneas que quedan detrás del respaldo y la parte oculta de las copias traseras."],
       ["6", "F5 (Top), ELLIPSE (I), LAYER",
        "Centro del agujero: desde el punto medio de la arista delantera superior de la "
        "base, 20 hacia 30°. Isocírculo de diámetro 10. Dibuje los ejes, pase todo a sus "
        "capas y apague el modo con ISODRAFT, Orthographic."]],
      [0.5, 1.75, 4.25])
fig("isocad_04_construccion.png", 6.4,
    "Figura 4. Construcción progresiva. Trazo grueso: lo nuevo de cada paso. Azul: referencias.")

# ------------------------------------------------------------------ 7
d.h1("Cotas isométricas (opcional)")
d.body("Las medidas de fabricación van en las vistas. Si se pide acotar el isométrico, "
       "la cota debe parecer apoyada en la cara:")
tabla(["Paso", "Comando", "Qué se hace"],
      [["1", "DIMALIGNED", "Acote la arista. La cota queda paralela a ella."],
       ["2", "DIMEDIT, opción Oblique",
        "Seleccione la cota y escriba el ángulo del eje al que deben quedar paralelas las "
        "líneas de referencia: 30, -30 (equivale a 150) o 90."],
       ["3", "STYLE",
        "Cree dos estilos de texto con ángulo oblicuo (Oblique Angle) 30 y -30. Use el que "
        "deje las letras paralelas a las líneas de referencia."]],
      [0.5, 1.9, 4.1])
fig("isocad_05_cotas.png", 4.7,
    "Figura 5. Cota alineada sin corregir y cota isométrica con líneas de referencia oblicuas.")

# ------------------------------------------------------------------ 8
d.h1("Del papel al CAD")
tabla(["N°", "Acción"],
      [["1", "Liste las medidas de su dibujo a mano: largos por eje, centros, radios y diámetros."],
       ["2", "Elija el punto de inicio: el vértice inferior más cercano al observador."],
       ["3", "Active ISODRAFT y dibuje la caja envolvente de cada bloque con LINE."],
       ["4", "Complete las caras planas, cambiando de isoplano con F5."],
       ["5", "Ubique los centros con líneas auxiliares y dibuje los isocírculos."],
       ["6", "Copie hacia el fondo lo que tenga espesor y trace las tangentes."],
       ["7", "Recorte con TRIM y borre las líneas ocultas y auxiliares. El isométrico no lleva líneas ocultas."],
       ["8", "Pase cada objeto a su capa, revise grosores y desactive el modo isométrico."]],
      [0.5, 6.0])

# ------------------------------------------------------------------ 9
d.h1("Errores frecuentes")
tabla(["Error", "Cómo se ve", "Cómo evitarlo"],
      [["Usar CIRCLE", "Un círculo perfecto que no calza con la cara",
        "ELLIPSE, opción Isocircle"],
       ["Isoplano equivocado", "La elipse queda girada respecto de la cara",
        "Mire el cursor y pulse F5 antes de dibujar"],
       ["Medir sobre una inclinada", "La pieza no cierra o queda con otra proporción",
        "Mida sobre los ejes y una los extremos"],
       ["Dejar el modo isométrico activo", "Las vistas ortogonales salen torcidas",
        "ISODRAFT, Orthographic antes de dibujar vistas"],
       ["Líneas duplicadas", "Trazos más gruesos al imprimir y TRIM que no recorta",
        "Use Endpoint; limpie con OVERKILL"],
       ["Todo en la capa 0", "Un solo grosor y ejes continuos",
        "Asigne capas antes de entregar"]],
      [1.8, 2.5, 2.2])

# ------------------------------------------------------------------ 10
d.h1("Regla de oro y taller")
d.destacado(1, "Antes de cada trazo mire el isoplano activo: mida solo sobre los ejes y "
               "dibuje cada círculo como isocírculo en el isoplano de su cara.")
d.destacado(3, "Taller de la clase. Dibuje la pieza guía en isométrico 2D, a escala 1:1 y con "
               "las cuatro capas del apunte. Sin líneas ocultas y con los ejes de los dos "
               "agujeros. Las cotas son opcionales. Entregue el archivo DWG y un PDF en A4 "
               "con cajetín ISO 7200.")

d.h1("Fuentes")
src = [
    "Autodesk. About 2D Isometric Drawing. https://help.autodesk.com/cloudhelp/2024/ENU/AutoCAD-LT/files/GUID-37463F74-0B06-46E2-8791-6C5B852A069D.htm",
    "Autodesk. ISODRAFT (Command). https://help.autodesk.com/cloudhelp/2015/ENU/AutoCAD-LT-MAC/files/GUID-AE050A14-4887-4940-A0C1-A4F6FEEE21CA.htm",
    "Autodesk. ISOPLANE (Command). https://help.autodesk.com/cloudhelp/2022/ENU/AutoCAD-LT/files/GUID-9B1EEA63-BEC1-413E-B69F-541B5865F1A1.htm",
    "Autodesk. ELLIPSE (Command). https://help.autodesk.com/cloudhelp/2026/ENU/AutoCAD-LT/files/GUID-45F9C588-2BA9-414D-8AFD-FB6B448BF273.htm",
    "Autodesk. DIMEDIT (Command). https://help.autodesk.com/cloudhelp/2018/ENU/AutoCAD-LT/files/GUID-4C422870-32A1-457B-8D1E-BD4AF7967FF1.htm",
]
for s in src:
    p = d.doc.add_paragraph(s, style="Normal")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(1)
    for r in p.runs:
        r.font.size = Pt(8)

d.headers(lhead="Sistemas de Representación · PCI 1119", rhead="Isométrico en AutoCAD",
          lfoot="Ingeniería Civil Industrial · UCT", cfoot="Prof. Miguel Godoy Díaz")
d.save(OUT, pdf=True)
print(OUT)
