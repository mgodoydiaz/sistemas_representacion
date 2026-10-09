"""Arma Proyecto_2_Enunciado.docx y .pdf (enunciado + rúbrica + anexo de piezas).

Uso:  python3 piezas_proyecto2.py && python3 build_proyecto2.py
Requiere la plantilla del profesor (skill plantilla-miguel) y LibreOffice para el PDF.
"""
import os
import re
import sys

from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = os.path.dirname(os.path.abspath(__file__))
PLANTILLA = os.environ.get("PLANTILLA_MIGUEL", "/mnt/skills/plugins/plantilla-miguel/scripts")
sys.path.insert(0, PLANTILLA)
from plantilla import Doc  # noqa: E402

FIG = os.path.join(HERE, "figuras")
d = Doc()


# ------------------------------------------------------------------ ayudas
def rich(p, text):
    """Escribe text en el párrafo p; **así** queda en negrita."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if part:
            p.add_run(part).bold = bool(i % 2)
    return p


def body(text):
    return rich(d.doc.add_paragraph(style="Normal"), text)


def item(text, mark="•"):
    p = d.doc.add_paragraph(style="Normal")
    pf = p.paragraph_format
    pf.left_indent = Cm(1.0)
    pf.first_line_indent = Cm(-0.5)
    pf.space_after = Pt(2)
    rich(p, f"{mark}\t{text}")
    pf.tab_stops.add_tab_stop(Cm(1.0))
    return p


def table(headers, rows, widths, size=9.5, bold_first=True):
    d.table(headers, rows)
    t = d.doc.tables[-1]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(int(w * 567)))
    for r_i, row in enumerate(t.rows):
        trPr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        trPr.append(cs)
        for c_i, cell in enumerate(row.cells):
            cell.width = Cm(widths[c_i])
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.space_before = Pt(1)
                for r in p.runs:
                    r.font.size = Pt(size)
                    if bold_first and c_i == 0 and r_i > 0:
                        r.font.bold = True
    d.doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def new_page(paragraph):
    paragraph.paragraph_format.page_break_before = True


# ------------------------------------------------------------------ contenido
d.title("Proyecto 2: Piezas de una válvula",
        "De las vistas al isométrico y del isométrico al plano, en AutoCAD · Semestre 2026-2")

body("En este proyecto usted trabajará íntegramente en AutoCAD con cuatro piezas inspiradas "
     "en los componentes de una válvula: bonete, prensaestopas, yugo y cuerpo. Las piezas "
     "están simplificadas con fines docentes. Todas las medidas están en milímetros.")
body("El proyecto equivale al **35 % de la nota final** y tiene dos partes. A usted se le "
     "asignan dos piezas distintas: una para la Parte A y otra para la Parte B.")

table(["Parte", "Usted recibe", "Usted entrega", "Peso"],
      [["A", "Tres vistas acotadas de una pieza (ISO-E)", "Su isométrico 2D", "40 %"],
       ["B", "El isométrico acotado de otra pieza",
        "El plano de la pieza: vistas con corte, cotas e isométrico propio", "60 %"]],
      [1.5, 6.0, 7.0, 1.6])

table(["Dato", "Valor"],
      [["Modalidad", "Individual o en parejas [por confirmar]"],
       ["Lanzamiento", "Clase 7, fecha [por definir]"],
       ["Entrega", "Clase 10, fecha [por definir]"],
       ["Variante Parte A", "V ____ (láminas «vistas» del anexo)"],
       ["Variante Parte B", "V ____ (láminas «isométrico» del anexo)"]],
      [4.0, 12.1])

d.h1("Parte A: de las vistas al isométrico (40 %)")
body("Con las tres vistas acotadas de su pieza, trace el isométrico 2D en AutoCAD.")
item("Dibuje en el espacio modelo a escala 1:1, con ISODRAFT activo. Las medidas de las "
     "vistas se llevan sin reducción sobre las tres direcciones isométricas.", "1.")
item("Trace círculos y arcos como isocírculos (comando ELIPSE, opción Isocírculo) en el "
     "isoplano que corresponda. Cambie de isoplano con F5.", "2.")
item("Cierre los cilindros con sus generatrices de contorno, tangentes a las elipses.", "3.")
item("Dibuje solo las aristas visibles. Recorte los tramos tapados y borre las líneas "
     "auxiliares. No deje líneas duplicadas ni cabos sueltos.", "4.")
item("No acote el isométrico. Use las capas indicadas en la sección 3.", "5.")
item("Presente el dibujo en una lámina con cajetín y genere el PDF.", "6.")

d.h1("Parte B: del isométrico al plano (60 %)")
body("Con el isométrico acotado de su pieza, dibuje el plano completo en una sola lámina.")
item("**Vistas.** Elija el alzado y dibuje solo las vistas necesarias, en ISO-E (primer "
     "diedro), alineadas y en correspondencia.", "1.")
item("**Corte.** Incluya al menos un corte total o un semicorte que reemplace las aristas "
     "ocultas. Indíquelo como en el capítulo 5 del apunte: traza de trazo y punto con "
     "extremos gruesos, flechas, letras y rótulo A-A sobre la vista cortada. En el semicorte "
     "las dos mitades se separan con el eje. En la vista cortada no se dibujan ocultas.", "2.")
item("**Rayado.** Líneas finas a 45°, con la misma separación en toda la pieza. Se raya solo "
     "el material que toca el plano. Los nervios cortados a lo largo no se rayan.", "3.")
item("**Acotado.** Completo y sin sobreacotar. Cada medida aparece una sola vez, en la vista "
     "donde mejor se lee. Use Ø y R. Las notas del isométrico entregado (profundidades, "
     "agujeros pasantes) deben quedar resueltas en su plano.", "4.")
item("**Isométrico propio.** Agregue en la misma lámina el isométrico 2D de la pieza, "
     "dibujado por usted, sin cotas y solo con aristas visibles.", "5.")
item("**Lámina.** Vistas e isométrico en una presentación con cajetín.", "6.")

d.h1("Requisitos de formato")
table(["Capa", "Contenido", "Tipo de línea", "Grosor"],
      [["Visible", "Aristas y contornos visibles", "Continua", "0,50 mm"],
       ["Oculta", "Aristas ocultas (solo si hacen falta)", "Trazos", "0,25 mm"],
       ["Ejes", "Ejes de simetría, centros de agujeros, trazas de corte", "Trazo y punto",
        "0,25 mm"],
       ["Cotas", "Cotas, líneas de referencia y notas", "Continua", "0,25 mm"],
       ["Rayado", "Rayado de las zonas cortadas", "Continua", "0,25 mm"],
       ["Cajetín", "Marco, cajetín y sus textos", "Continua", "0,50 y 0,25 mm"]],
      [2.4, 7.6, 3.4, 2.7])
item("Tipos de línea y grosores según ISO 128, asignados por capa y no objeto por objeto.")
item("Dibujo a escala 1:1 en el espacio modelo. Marco y cajetín en la presentación.")
item("Ventana con escala normalizada: 1:1, 1:2 o 2:1. La escala del cajetín debe coincidir "
     "con la de la ventana.")
item("Cajetín según ISO 7200 en la esquina inferior derecha: título, nombre, fecha, escala, "
     "formato, método de proyección y número de lámina.")
item("Formato de hoja: A4 para la Parte A y A3 para la Parte B [por confirmar].")

d.h1("Entrega")
body("Por cada parte se entregan dos archivos: el DWG y el PDF de la lámina. En total son "
     "cuatro archivos, con estos nombres:")
d.code("PCI1119_P2_Apellido_Nombre_ParteA.dwg\nPCI1119_P2_Apellido_Nombre_ParteA.pdf\n"
       "PCI1119_P2_Apellido_Nombre_ParteB.dwg\nPCI1119_P2_Apellido_Nombre_ParteB.pdf",
       highlight=False, size=9)
item("Medio de entrega: [por definir].")
item("El DWG debe abrir sin errores y sin pedir archivos externos.")
item("El PDF debe mostrar los grosores de línea. Revíselo antes de enviar.")
item("Si trabaja en pareja [por confirmar], use los apellidos de ambos integrantes.")

d.h1("Bonificación opcional")
body("Puede sumar hasta 5 décimas a la nota del proyecto [por confirmar] con una de estas "
     "opciones, aplicada a la pieza de la Parte B:")
item("Vista auxiliar de una cara plana inclinada, en verdadera magnitud. Aplica a las "
     "piezas que tienen una: V1 (cara del nervio) y V3 (chaflán del puente).")
item("Generar una de las vistas con Python y la biblioteca ezdxf. Se entrega el script y "
     "el DXF que produce.")

h = d.doc.add_paragraph("Rúbrica", style="Heading 1")
new_page(h)
body("Cada parte se evalúa sobre 100 puntos. En cada criterio se asigna uno de cuatro "
     "niveles. El puntaje del proyecto es 0,4 × Parte A + 0,6 × Parte B.")

HDR = ["Criterio", "Logrado", "Medianamente logrado", "Por lograr", "No observado"]
W = [3.0, 3.6, 3.5, 3.5, 2.6]


def rubrica(titulo, filas, salto=False):
    h2 = d.doc.add_paragraph(titulo, style="Heading 2")
    if salto:
        new_page(h2)
    rows = []
    for nombre, pts, niveles in filas:
        rows.append([f"{nombre} ({pts[0]} pts)"] +
                    [f"{p} pts. {t}" for p, t in zip(pts, niveles)])
    assert sum(f[1][0] for f in filas) == 100
    table(HDR, rows, W, size=8.5)


rubrica("Parte A: isométrico (100 puntos)", [
    ("Isoplanos y direcciones", (20, 12, 6, 0),
     ["Todas las aristas siguen las direcciones isométricas y cada cara está en su isoplano.",
      "Uno o dos tramos fuera de dirección. La pieza se reconoce.",
      "Errores repetidos de dirección que deforman la pieza.",
      "No es un isométrico o no se entrega."]),
    ("Medidas fieles a las vistas", (20, 12, 6, 0),
     ["Todas las medidas coinciden con las vistas, a escala 1:1.",
      "Hasta dos medidas erróneas o un elemento mal ubicado.",
      "Varias medidas erróneas o proporciones alteradas.",
      "Medidas sin relación con las vistas."]),
    ("Isocírculos y arcos", (20, 12, 6, 0),
     ["Isocírculos en el isoplano correcto, tangencias limpias y generatrices de contorno.",
      "Un isocírculo en el isoplano equivocado o tangencias con pequeños saltos.",
      "Varios isocírculos mal orientados o círculos sin proyectar.",
      "Sin curvas o sin relación con la pieza."]),
    ("Aristas visibles y trazo limpio", (15, 9, 5, 0),
     ["Solo aristas visibles. Sin duplicados, cabos sueltos ni tramos sin recortar.",
      "Alguna arista oculta dibujada o dos o tres cabos sueltos.",
      "Muchas líneas sobrantes, duplicadas o sin recortar.",
      "Dibujo de alambre, sin depurar."]),
    ("Capas y grosores", (10, 6, 3, 0),
     ["Cada objeto en su capa. Grosor y tipo de línea por capa.",
      "Algunos objetos en la capa equivocada o propiedades forzadas a mano.",
      "Capas creadas, pero casi todo en una sola.",
      "Todo en la capa 0."]),
    ("Lámina, cajetín y PDF", (15, 9, 5, 0),
     ["Ventana a escala normalizada, cajetín completo y PDF legible con grosores.",
      "Falta un dato del cajetín o la escala indicada no coincide con la ventana.",
      "Escala no normalizada, cajetín incompleto o PDF sin grosores.",
      "Sin presentación o sin PDF."]),
])

rubrica("Parte B: plano de la pieza (100 puntos)", [
    ("Elección y disposición de vistas", (15, 9, 5, 0),
     ["Vistas necesarias y suficientes, en ISO-E, alineadas y en correspondencia.",
      "Sobra una vista o falta un detalle. Pequeños desalineamientos.",
      "Disposición que no es ISO-E o vistas que no calzan entre sí.",
      "Una sola vista o vistas sin relación."]),
    ("Corte", (20, 12, 6, 0),
     ["Traza, flechas, letras y rótulo correctos. Rayado solo en el material cortado. Sin "
      "ocultas. Nervios sin rayar.",
      "Corte bien elegido con una omisión: falta el rótulo, una flecha o hay una zona mal "
      "rayada.",
      "Corte mal ubicado, rayado sobre huecos o con ocultas que repiten el corte.",
      "Sin corte."]),
    ("Acotado", (20, 12, 6, 0),
     ["Completo y sin repeticiones. Ø y R bien indicados. Cotas en su capa y con estilo "
      "propio.",
      "Faltan hasta dos cotas o hay una repetida. Estilo correcto.",
      "Faltan varias cotas, hay sobreacotado o cotas sobre aristas ocultas.",
      "Sin cotas."]),
    ("Tipos de línea, capas y grosores", (10, 6, 3, 0),
     ["Visibles gruesas. Ocultas, ejes, cotas y rayado finos. Todo en su capa.",
      "Algún tipo de línea o grosor equivocado. Falta el eje de algún agujero.",
      "Tipos de línea forzados a mano o sin diferencia de grosores.",
      "Una sola capa y un solo tipo de línea."]),
    ("Isométrico propio", (15, 9, 5, 0),
     ["Coherente con las vistas, con isocírculos correctos y solo aristas visibles.",
      "Reconocible, con uno o dos errores de medida o de isoplano.",
      "Incompleto o con errores que cambian la forma.",
      "No se entrega."]),
    ("Lámina", (15, 9, 5, 0),
     ["Buena distribución, escala de ventana normalizada, cajetín ISO 7200 completo y PDF "
      "legible.",
      "Distribución apretada o falta un dato del cajetín.",
      "Escala no normalizada, cajetín incompleto o PDF ilegible.",
      "Sin presentación."]),
    ("Formalidad de entrega", (5, 3, 2, 0),
     ["Nombres de archivo correctos. El DWG abre sin errores.",
      "Nombres con errores menores.",
      "Falta un archivo o el DWG pide archivos externos.",
      "El DWG no abre o no se entrega."]),
], salto=True)

# ------------------------------------------------------------------ anexo
h = d.doc.add_paragraph("Anexo: láminas de las piezas", style="Heading 1")
new_page(h)
body("Cada pieza tiene dos láminas. La lámina de vistas se usa en la Parte A y la lámina "
     "de isométrico en la Parte B. Las figuras no están a escala: trabaje con las cotas. "
     "Todas las piezas son simétricas respecto de sus dos planos verticales medios.")
table(["Variante", "Pieza", "Función en la válvula", "Envolvente (mm)"],
      [["V1", "Bonete", "Tapa del cuerpo. Guía el vástago y aloja la empaquetadura.",
        "80 × 80 × 60"],
       ["V2", "Prensaestopas", "Comprime la empaquetadura alrededor del vástago.",
        "110 × 40 × 50"],
       ["V3", "Yugo", "Puente que sostiene la tuerca del vástago.", "110 × 40 × 84"],
       ["V4", "Cuerpo", "Bloque con el paso del fluido y el alojamiento de la compuerta.",
        "90 × 50 × 60"]],
      [1.9, 2.8, 8.2, 3.2])

PIEZAS = [("V1", "Bonete"), ("V2", "Prensaestopas"), ("V3", "Yugo"), ("V4", "Cuerpo")]
n = 0
for cod, nombre in PIEZAS:
    for suf, desc, parte in (("vistas", "vistas acotadas", "Parte A"),
                             ("iso", "isométrico acotado", "Parte B")):
        n += 1
        h = d.doc.add_paragraph(f"{cod} {nombre}: {desc} ({parte})", style="Heading 2")
        new_page(h)
        d.image(os.path.join(FIG, f"{cod}_{suf}.png"), width_in=6.3,
                caption=f"Figura {n}. {cod} {nombre}, {desc}.")

d.headers(lhead="Sistemas de Representación · PCI 1119",
          rhead="Proyecto 2 · Piezas de una válvula",
          lfoot="Ingeniería Civil Industrial · UCT", cfoot="Prof. Miguel Godoy Díaz")
out = os.path.join(HERE, "Proyecto_2_Enunciado.docx")
d.save(out, pdf=True)
print(out)
