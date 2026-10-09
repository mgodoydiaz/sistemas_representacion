"""Construye Apuntes_Unidad1_PCI1119_v3 (docx + pdf) a partir de la v2.

- Carga ref/Apuntes_Unidad1_PCI1119_v2.docx (no se modifica).
- Reemplaza la sección 2.2 "El cajetín" por la versión ampliada.
- Ajusta el Taller 2.8 ("seis campos exigidos en el curso").
- Agrega los capítulos 4 (Vistas auxiliares) y 5 (Cortes y secciones).
- Figuras: figuras/apunte_*.png (generadas por figuras_apunte_v3.py).
Uso: python3 build_apunte_v3.py
"""
import copy
import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "ref", "Apuntes_Unidad1_PCI1119_v2.docx")
OUT_DIR = os.path.join(HERE, "salida")
OUT = os.path.join(OUT_DIR, "Apuntes_Unidad1_PCI1119_v3.docx")
FIG = os.path.join(HERE, "figuras")

TITLE = "Apuntes Unidad 1 · Fundamentos de la representación"
RHEAD = "Unidad 1 · Fundamentos de la representación"
TEXT_W = 9360  # ancho útil en dxa (6,5 in)


# ====================================================================== helpers
class Writer:
    """Agrega bloques al final del cuerpo con los estilos del molde de Miguel,
    imitando el formato de la v2 (listas con viñeta manual, tablas con encabezado
    repetido, imágenes con keepNext)."""

    def __init__(self, doc):
        self.doc = doc
        self.body = doc.element.body

    # --- texto con **negrita**
    def _runs(self, p, text):
        parts = re.split(r"(\*\*.+?\*\*)", text)
        for part in parts:
            if not part:
                continue
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2])
                r.bold = True
            else:
                r = p.add_run(part)
                r.bold = False
        return p

    def para(self, text, style="Normal"):
        p = self.doc.add_paragraph(style=style)
        return self._runs(p, text)

    def h1(self, text, page_break=True):
        p = self.doc.add_paragraph(text, style="Heading 1")
        p.paragraph_format.page_break_before = page_break
        return p

    def h2(self, text):
        return self.doc.add_paragraph(text, style="Heading 2")

    def h3(self, text):
        return self.doc.add_paragraph(text, style="Heading 3")

    def dest(self, n, text):
        return self.para(text, f"Destacado {n}")

    def _list_item(self, marker, text):
        p = self.doc.add_paragraph(style="List Paragraph")
        pPr = p._p.get_or_add_pPr()
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "left")
        tab.set(qn("w:pos"), "504")
        tabs.append(tab)
        pPr.append(tabs)
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:after"), "40")
        pPr.append(sp)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "504")
        ind.set(qn("w:hanging"), "288")
        pPr.append(ind)
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "both")
        pPr.append(jc)
        r = p.add_run(marker)
        r.add_tab()
        self._runs(p, text)
        return p

    def bullets(self, items):
        for it in items:
            self._list_item("•", it)

    def numbered(self, items, keep=False):
        for i, it in enumerate(items, 1):
            p = self._list_item(f"{i}.", it)
            if keep:
                p.paragraph_format.keep_with_next = True

    def table(self, headers, rows, widths=None):
        n = len(headers)
        widths = widths or [1.0 / n] * n
        tw = [int(TEXT_W * w / sum(widths)) for w in widths]
        t = self.doc.add_table(rows=1, cols=n)
        t.style = "Table"
        # grid
        grid = t._tbl.tblGrid
        for gc, w in zip(grid.findall(qn("w:gridCol")), tw):
            gc.set(qn("w:w"), str(w))
        all_rows = [headers] + rows
        for ri, rdata in enumerate(all_rows):
            row = t.rows[0] if ri == 0 else t.add_row()
            trPr = row._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:cantSplit"))
            if ri == 0:
                trPr.append(OxmlElement("w:tblHeader"))
            for ci, val in enumerate(rdata):
                cell = row.cells[ci]
                tcPr = cell._tc.get_or_add_tcPr()
                tcW = tcPr.find(qn("w:tcW"))
                if tcW is None:
                    tcW = OxmlElement("w:tcW")
                    tcPr.insert(0, tcW)
                tcW.set(qn("w:type"), "dxa")
                tcW.set(qn("w:w"), str(tw[ci]))
                p = cell.paragraphs[0]
                if ri < len(all_rows) - 1:
                    p.paragraph_format.keep_with_next = True
                if ri == 0:
                    p.add_run(str(val))
                else:
                    self._runs(p, str(val))
        return t

    def image(self, name, width_in, caption):
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf = p.paragraph_format
        pf.keep_with_next = True
        pf.space_before = Pt(8)
        p.add_run().add_picture(os.path.join(FIG, name), width=Inches(width_in))
        c = self.doc.add_paragraph(caption, style="Caption")
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return p

    def taller_lead(self, label, text=""):
        p = self.para(f"**{label}**" + (f" {text}" if text else ""))
        p.paragraph_format.keep_with_next = True
        return p


def body_children(doc):
    return [c for c in doc.element.body if not c.tag.endswith("}sectPr")]


def style_name(doc, el):
    from docx.text.paragraph import Paragraph
    if not el.tag.endswith("}p"):
        return None
    return Paragraph(el, doc).style.name


def text_of(doc, el):
    from docx.text.paragraph import Paragraph
    return Paragraph(el, doc).text if el.tag.endswith("}p") else ""


# ====================================================================== contenido
def seccion_cajetin(w):
    """Sección 2.2 El cajetín (versión ampliada). Devuelve bloques al final."""
    w.h2("El cajetín")
    w.dest(1, "El cajetín (o cuadro de rotulación) es el bloque de identificación del "
              "plano. Va en la esquina inferior derecha, apoyado en el recuadro, y se lee "
              "en la misma dirección que el dibujo.")
    w.para("Es lo primero que lee quien recibe el plano: sin él, la lámina no es un "
           "documento técnico. En A4 vertical ocupa todo el ancho útil del recuadro: "
           "**180 mm** (210 − 20 − 10).")
    w.image("apunte_02_lamina_a4.png", 2.5,
            "Figura 2. Lámina A4 vertical: recuadro ISO 5457 y cajetín de 180 mm.")
    w.para("**Norma.** La norma internacional vigente es la **ISO 7200:2004**. No fija "
           "un dibujo: define los campos de datos del cajetín. Declara **ocho campos "
           "obligatorios** y un ancho de **180 mm**, igual en todos los formatos. En Chile "
           "rige la **NCh 14.Of93** (Cuadro de rotulación), basada en la edición anterior, "
           "ISO 7200:1984: exige número del dibujo, título y propietario, y limita el largo "
           "a **170 mm**. La acompaña la **NCh 13.Of93**, que fija formatos y márgenes. El "
           "curso adopta la ISO 7200:2004 y ubica los datos indicativos (escala, método de "
           "proyección, unidades) como indica la NCh 14: en una zona suplementaria.")
    w.image("apunte_03_cajetin_iso7200.png", 6.3,
            "Figura 3. Cajetín ISO 7200:2004 con ejemplo de llenado para PCI 1119.")
    w.para("Los seis **campos exigidos en este curso** se corresponden así con la norma:")
    w.table(["Campo exigido en el curso", "Campo en ISO 7200:2004", "Dónde va"], [
        ["Título", "Título (obligatorio)", "Zona de identificación"],
        ["Autor", "Creado por (obligatorio)", "Cajetín"],
        ["Fecha", "Fecha de edición (obligatorio), AAAA-MM-DD", "Cajetín"],
        ["Número de lámina", "Número de hoja (obligatorio)", "Zona de identificación"],
        ["Escala", "No es campo de la norma", "ISO: fuera del cajetín. NCh 14: zona suplementaria"],
        ["Símbolo de diedro", "No es campo de la norma", "ISO: fuera del cajetín. NCh 14: zona suplementaria"],
    ], widths=[1.1, 1.5, 1.6])
    w.para("Los otros cuatro obligatorios de la ISO (propietario legal, número de "
           "identificación, aprobado por y tipo de documento) se llenan como en la Figura 3.")
    w.dest(2, "**Por qué el curso los pide dentro.** La ISO 7200:2004 deja la escala y el "
              "símbolo de proyección fuera del cajetín, porque cambian de un plano a otro. "
              "En este curso se exigen dentro, en la fila superior, por simplicidad: son lo "
              "primero que se corrige. La escala se escribe con dos puntos (1:2), nunca 1/2 "
              "ni 50 %.")
    w.dest(3, "Regla de oro: el cajetín tiene el mismo tamaño en todas las láminas; lo que "
              "cambia es la escala declarada en él.")
    w.h3("El cajetín en AutoCAD")
    w.dest(2, "**Idea clave:** la pieza se dibuja a 1:1 en el espacio modelo; el cajetín "
              "vive en la presentación a escala 1; la escala del dibujo se fija en la ventana.")
    w.numbered([
        "**Presentación.** Clic en la pestaña Presentación (Layout). Lo que sigue se dibuja "
        "en espacio papel, en milímetros reales.",
        "**Página.** PAGESETUP (PREPPAGINA): impresora DWG To PDF.pc3, papel ISO A4 o A3, "
        "escala 1:1.",
        "**Marco.** Recuadro ISO 5457: 20 mm a la izquierda y 10 mm en los demás lados.",
        "**Bloque.** Dibuje las celdas del cajetín (180 mm), defina los datos variables con "
        "ATTDEF (ATRDEF) y agrupe todo con BLOCK (BLOQUE), punto base en la esquina inferior "
        "derecha.",
        "**Insertar.** INSERT a escala 1 en la esquina del recuadro; el programa pide cada "
        "atributo. Para corregir: doble clic, EATTEDIT (EDITATR).",
        "**Campos (opcional).** FIELD (CAMPO) llena solo la fecha o la escala de la ventana.",
        "**Ventana.** MVIEW (VMULT) dentro del recuadro, sin tapar el cajetín. Fije una "
        "escala normalizada (1:2, 1:5) y bloquee la visualización.",
        "**Imprimir.** PLOT (TRAZAR) a PDF: área Presentación, escala 1:1, estilo "
        "monochrome.ctb.",
    ])
    w.para("La plantilla **acadiso.dwt** trae unidades y cotas ISO, pero no trae cajetín. "
           "Se hace una vez y se guarda como plantilla propia (.dwt).")


def capitulo_4(w):
    w.h1("Vistas auxiliares")
    w.para("Las vistas principales bastan cuando las caras de la pieza son paralelas a los "
           "planos de proyección. Si una cara está inclinada, ninguna de ellas la muestra "
           "en su forma real.")

    w.h2("El problema: la cara inclinada se deforma")
    w.para("El bloque de la Figura 7 tiene un chaflán a 45° con un agujero Ø12. En el "
           "alzado, la cara inclinada se ve como una línea. En la planta aparece, pero "
           "deformada: mide **30 mm en vez de 42,4 mm** y el agujero se ve como **elipse**. "
           "Sobre esa vista no se puede medir ni acotar.")
    w.image("apunte_07_por_que_auxiliar.png", 5.4,
            "Figura 7. La cara inclinada no se ve en verdadera magnitud en ninguna vista principal.")

    w.h2("Definición")
    w.dest(1, "La vista auxiliar es la proyección ortogonal de la pieza sobre un plano "
              "paralelo a una cara inclinada. Muestra esa cara en verdadera magnitud (VM).")
    w.para("**Condición de uso:** la cara debe verse **de canto** (como una línea) en una "
           "vista principal. Desde esa vista se proyecta.")
    w.dest(1, "La línea de referencia (LR) marca el plano desde el que se miden las "
              "profundidades. Va en las dos vistas: paralela a la arista de canto en la "
              "auxiliar y sobre la cara trasera en la planta.")

    w.h2("Procedimiento")
    w.table(["Paso", "Qué se hace", "Qué se mide"], [
        ["1", "Ubicar la vista donde la cara está de canto (aquí, el alzado).",
         "Nada: se identifica la arista."],
        ["2", "Trazar la LR paralela a la arista de canto. En la planta, la LR es la cara trasera.",
         "Distancia libre, 15 a 25 mm."],
        ["3", "Trazar proyectantes finas perpendiculares a la arista, desde cada vértice y "
              "desde el centro del agujero.", "Nada: escuadra apoyada en la arista."],
        ["4", "Llevar cada profundidad desde la planta, medida desde la LR, y unir los puntos.",
         "La medida que no se ve en la vista de canto: la profundidad (40)."],
    ], widths=[0.5, 3.0, 1.9])
    w.image("apunte_08_auxiliar_pasos.png", 4.7,
            "Figura 8. Vista auxiliar en cuatro pasos (los números coinciden con la tabla).")
    w.dest(2, "**Truco:** la profundidad se mide siempre desde la línea de referencia, en "
              "las dos vistas. Es la regla del Capítulo 3: se mide una vez y se traslada.")

    w.h2("Vista auxiliar parcial y rotulado")
    w.para("Se dibuja **solo la cara inclinada**: el resto de la pieza saldría deformado y "
           "no aporta. La vista se limita con una línea fina a mano alzada o en zigzag "
           "(Figura 8, paso 4). Es la forma habitual en los planos.")
    w.para("Si la vista queda alineada con su vista de origen, no lleva rótulo. Si se "
           "desplaza a otro lugar de la lámina, se indica con una **flecha y una letra "
           "mayúscula** junto a la vista de origen, y la misma letra sobre la vista auxiliar "
           "(ISO 128-3).")

    w.h2("Ahora usted")
    w.para("Mismo bloque, con el chaflán a 30°. Complete los pasos 3 y 4.")
    w.table(["Paso", "Resolución"], [
        ["1", "La arista de canto es la línea del chaflán en el alzado."],
        ["2", "LR paralela a esa línea, a 20 mm; en la planta, sobre la cara trasera."],
        ["3", "Complete: ¿desde qué puntos y en qué dirección van las proyectantes?"],
        ["4", "Complete: ¿qué medida se lleva y desde dónde se mide?"],
    ], widths=[0.5, 5.0])

    w.h2("Errores frecuentes")
    w.table(["Error", "Cómo se ve", "Cómo evitarlo"], [
        ["LR no paralela a la arista de canto", "La cara sale deformada otra vez",
         "Marcar primero la arista (paso 1)"],
        ["Proyectantes no perpendiculares", "Vista corrida o torcida",
         "Escuadra apoyada en la arista de canto"],
        ["Transferir la medida equivocada", "Se lleva el ancho o el alto",
         "Llevar la medida que no se ve de canto"],
        ["Proyectar desde una vista sin la cara de canto", "No se obtiene VM",
         "Revisar la condición de uso"],
        ["Dibujar la auxiliar completa", "Pieza deformada, lámina saturada",
         "Usar la vista parcial"],
    ], widths=[1.6, 1.5, 1.6])
    w.dest(3, "Regla de oro: una cara solo se ve en verdadera magnitud cuando se mira "
              "perpendicularmente a ella: primero búsquela de canto.")
    w.para("Este mismo procedimiento, aplicado a rectas y planos, es el **cambio de plano** "
           "que se usará para hallar verdaderas magnitudes en el sistema diédrico.")

    w.h2("Taller de la clase")
    w.taller_lead("Encargo.", "Dada una pieza en isometría con una cara inclinada a 30° y "
                  "una ranura, dibujar en A4 con cajetín:")
    w.numbered(keep=True, items=["Alzado y planta en primer diedro.",
                "Vista auxiliar parcial de la cara inclinada, con su línea de referencia y "
                "proyectantes visibles en trazo fino."])
    w.taller_lead("Criterios de corrección.")
    w.table(["Criterio", "Qué se evalúa"], [
        ["Construcción", "LR paralela a la arista de canto y proyectantes perpendiculares"],
        ["Verdadera magnitud", "Medidas de la cara y de la ranura sin deformar"],
        ["Vista parcial", "Solo la cara inclinada, limitada por línea de rotura fina"],
    ])


def capitulo_5(w):
    w.h1("Cortes y secciones")

    w.h2("Por qué cortar")
    w.para("Una pieza con huecos interiores llena sus vistas de aristas ocultas, y con "
           "muchas ocultas la vista se vuelve ilegible. La solución: cortar la pieza con un "
           "plano imaginario, retirar la parte que queda entre el observador y el plano, y "
           "dibujar lo que queda (Figura 9).")
    w.image("apunte_09_que_es_cortar.png", 6.3,
            "Figura 9. Del plano de corte a la vista cortada.")

    w.h2("Corte, sección y rotura")
    w.dest(1, "Corte: representa lo que toca el plano de corte más lo que se ve detrás de él.")
    w.dest(1, "Sección: representa solo la figura que el plano deja sobre la pieza.")
    w.para("La **rotura** es otra cosa: una interrupción convencional de la pieza para "
           "ahorrar espacio o mostrar un detalle interior. No es un corte.")

    w.h2("Cómo se indica el corte")
    w.para("Según **ISO 128-3:2022**:")
    w.numbered([
        "**Traza del plano:** trazo y punto fino, con tramos gruesos en los extremos y en "
        "los cambios de dirección.",
        "**Flechas** en los extremos, perpendiculares a la traza, apuntando hacia donde se mira.",
        "**Letra mayúscula** junto a cada flecha, la misma en ambas.",
        "**Rótulo** A-A encima de la vista cortada. Si hay varios cortes: B-B, C-C.",
    ])
    w.para("En la vista cortada **no se dibujan aristas ocultas**.")
    w.image("apunte_10_indicacion_AA.png", 3.8,
            "Figura 10. Indicación del corte A-A en la planta y vista cortada rotulada.")

    w.h2("Rayado")
    w.para("El rayado dice dónde había material. Reglas (ISO 128-3, cláusula 7):")
    w.bullets([
        "Líneas **finas y paralelas**, a **45°** del contorno principal (30° o 60° si el "
        "contorno está a 45°).",
        "**Separación constante**, proporcional al área. En A4 se usa habitualmente 2 a 3 mm "
        "(práctica común; la norma pide proporcionalidad y un mínimo de 0,7 mm).",
        "Se raya **solo el material** que toca el plano. Los huecos quedan en blanco.",
        "**Misma pieza, mismo rayado** en todas sus zonas y vistas. Piezas vecinas: otra "
        "dirección u otra separación.",
        "El rayado **llega al contorno sin cruzarlo** y se interrumpe para inscribir cotas.",
    ])
    w.image("apunte_11_rayado.png", 3.5,
            "Figura 11. Rayado correcto e incorrecto de una placa cortada por su agujero.")

    w.h2("Tipos de corte")
    w.table(["Tipo", "Cuándo se usa", "Cómo se indica"], [
        ["Corte total", "Ver todo el interior con un solo plano", "Traza completa con letras A-A"],
        ["Semicorte", "Pieza simétrica: mitad en vista, mitad cortada",
         "Mitades separadas por **eje**, nunca por línea gruesa; sin ocultas en la mitad en vista"],
        ["Corte parcial", "Un detalle interior pequeño (chavetero, agujero)",
         "Límite con línea fina a mano alzada o zigzag; sin letras"],
        ["Por planos paralelos", "Detalles interiores no alineados",
         "Traza quebrada, gruesa en los quiebres; en la vista no se dibujan los cambios de plano"],
        ["Alineado", "Pieza de revolución con detalles a distinto ángulo",
         "Dos planos que se cortan; uno se gira hasta el otro antes de proyectar"],
    ], widths=[1.1, 1.6, 2.2])
    w.image("apunte_12_total_vs_semicorte.png", 4.0,
            "Figura 12. Corte total y semicorte de la misma pieza, con su traza en la planta.")

    w.h2("Secciones abatida y desplazada")
    w.table(["Sección", "Dónde se dibuja", "Contorno", "Rótulo"], [
        ["Abatida", "Sobre la vista, girada 90° en su lugar", "Fino", "No lleva"],
        ["Desplazada", "Fuera de la vista, cerca de ella", "Grueso", "B-B (o unida por un eje)"],
    ], widths=[1.0, 2.2, 0.8, 1.5])
    w.image("apunte_13_parcial_secciones.png", 4.9,
            "Figura 13. Árbol con chavetero: corte parcial, sección abatida y sección desplazada.")

    w.h2("Lo que no se corta a lo largo")
    w.dest(3, "No se cortan longitudinalmente: ejes y árboles macizos, chavetas, pernos, "
              "tornillos, tuercas, pasadores, remaches, nervios y radios de ruedas.")
    w.para("Son macizos: rayarlos no informa y confunde. Un eje rayado parece fundido con "
           "la pieza que lo rodea. **Transversalmente sí se cortan** (Figura 13, B-B).")
    w.image("apunte_14_eje_en_conjunto.png", 4.5,
            "Figura 14. Conjunto cortado: el árbol y la chaveta quedan sin rayar.")

    w.h2("Cómo se indica un eje cortado")
    w.para("La frase admite tres situaciones distintas:")
    w.bullets([
        "**(a) Árbol seccionado.** No se hace corte longitudinal. Su perfil (chavetero, "
        "planos, agujeros) se muestra con una **sección abatida o desplazada** (Figura 13). "
        "Si el detalle es pequeño, basta un corte parcial.",
        "**(b) Eje largo interrumpido.** Se dibujan los extremos y se acercan. ISO 128-3 "
        "marca el límite con **línea continua fina a mano alzada o en zigzag**. La rotura en "
        "«S» es una convención tradicional, no exigida por la ISO vigente, pero aún frecuente "
        "en talleres. La cota indica la **longitud real** y su línea **no se interrumpe** "
        "(ISO 129-1).",
        "**(c) Pieza simétrica dibujada a medias.** Se dibuja hasta el eje y se marca el "
        "**símbolo de simetría**: dos trazos cortos, finos y paralelos, perpendiculares al "
        "eje en cada extremo. A diferencia del semicorte, la otra mitad no se dibuja.",
    ])
    w.image("apunte_15_roturas_simetria.png", 5.0,
            "Figura 15. Eje acortado con rotura (ISO y en «S») y placa simétrica dibujada a medias.")

    w.h2("Errores frecuentes")
    w.table(["Error", "Cómo se ve", "Cómo evitarlo"], [
        ["Dibujar solo la superficie cortada en un corte", "Faltan aristas detrás del plano",
         "Corte = plano + lo que se ve detrás"],
        ["Ocultas en la vista cortada", "Trazos dentro o fuera del rayado", "Omitirlas"],
        ["Rayar huecos", "Rayado dentro de agujeros", "Rayar solo donde había material"],
        ["Flechas invertidas o sin letras", "Vista cortada al lado equivocado",
         "La flecha dice hacia dónde se mira"],
        ["Línea gruesa entre mitades del semicorte", "Parece una arista real", "Separar con eje"],
        ["Rayar ejes o chavetas a lo largo", "Eje macizo lleno de rayado", "Revisar la lista cerrada"],
        ["Acotar la longitud dibujada", "Cota 120 en un eje de 600", "Acotar siempre la medida real"],
    ], widths=[1.7, 1.5, 1.5])
    w.dest(3, "Regla de oro: se raya solo el material que toca el plano de corte; la flecha "
              "indica hacia dónde se mira.")

    w.h2("Taller de la clase")
    w.taller_lead("Encargo.", "En A4 con cajetín:")
    w.numbered(keep=True, items=[
        "Buje con brida de la Figura 9 en **semicorte**, con la traza indicada en la planta.",
        "Árbol de la Figura 13 acortado con rotura, con una **sección desplazada B-B** en el "
        "chavetero y su longitud real acotada.",
    ])
    w.taller_lead("Criterios de corrección.")
    w.table(["Criterio", "Qué se evalúa"], [
        ["Indicación", "Traza con tramos gruesos, flechas en la dirección correcta, letras y rótulo"],
        ["Rayado", "Fino, a 45°, uniforme, solo en material cortado; eje sin rayar"],
        ["Limpieza de la vista cortada", "Sin ocultas; semicorte separado por eje; cota real en el eje con rotura"],
    ])


# ====================================================================== armado
def main():
    doc = Document(SRC)
    w = Writer(doc)
    body = doc.element.body

    # --- título y encabezado
    for p in doc.paragraphs:
        if p.style.name == "Title":
            for r in p.runs[1:]:
                r._r.getparent().remove(r._r)
            p.runs[0].text = TITLE
            break
    for sec in doc.sections:
        for p in sec.header.paragraphs:
            for r in p.runs:
                if r.text.startswith("Unidad 1"):
                    r.text = RHEAD

    # --- Taller 2.8: "seis campos obligatorios" -> "seis campos exigidos en el curso"
    n_fix = 0
    for p in doc.paragraphs:
        for r in p.runs:
            if "seis campos obligatorios" in r.text:
                r.text = r.text.replace("seis campos obligatorios",
                                        "seis campos exigidos en el curso")
                n_fix += 1
    assert n_fix == 1, n_fix

    # --- renumerar los pies heredados de la v2 que quedan después del cajetín
    #     (las dos figuras nuevas del 2.2 pasan a ser las Figuras 2 y 3)
    ren = {"Figura 2.": "Figura 4.", "Figura 3.": "Figura 5.", "Figura 4.": "Figura 6."}
    for p in doc.paragraphs:
        if p.style.name == "Caption" and p.runs:
            for old, nw in ren.items():
                if p.runs[0].text.startswith(old):
                    p.runs[0].text = p.runs[0].text.replace(old, nw, 1)
                    break

    # --- ubicar 2.2 El cajetín y el siguiente encabezado
    kids = body_children(doc)
    i0 = next(i for i, el in enumerate(kids)
              if style_name(doc, el) == "Heading 2" and text_of(doc, el).strip() == "El cajetín")
    i1 = next(i for i in range(i0 + 1, len(kids))
              if style_name(doc, kids[i]) in ("Heading 1", "Heading 2"))
    anchor = kids[i1]
    for el in kids[i0:i1]:
        body.remove(el)

    # --- generar el cajetín al final y moverlo antes del ancla
    n_before = len(body_children(doc))
    seccion_cajetin(w)
    new = body_children(doc)[n_before:]
    for el in new:
        anchor.addprevious(el)

    # --- capítulos nuevos al final
    capitulo_4(w)
    capitulo_5(w)

    # --- saltos de página: en vez de un párrafo vacío con salto (que en la v2 podía
    #     caer solo en una página y dejarla en blanco), el Heading 1 lleva
    #     "salto de página anterior". El contenido no cambia.
    kids = body_children(doc)
    h1_idx = [i for i, el in enumerate(kids) if style_name(doc, el) == "Heading 1"]
    for i in h1_idx[1:]:
        el = kids[i]
        if True:
            prev = kids[i - 1]
            brs = prev.findall(".//" + qn("w:br"))
            if prev.tag.endswith("}p") and brs and not text_of(doc, prev).strip() \
                    and all(b.get(qn("w:type")) == "page" for b in brs):
                body.remove(prev)
            from docx.text.paragraph import Paragraph
            Paragraph(el, doc).paragraph_format.page_break_before = True

    os.makedirs(OUT_DIR, exist_ok=True)
    doc.save(OUT)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", OUT_DIR, OUT],
                   check=False, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    print(OUT)


if __name__ == "__main__":
    main()
