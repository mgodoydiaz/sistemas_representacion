"""Genera los apuntes oficiales de PCI 1119 por módulo, a partir de los apuntes existentes.
Uso: python3 build_oficiales.py <raiz_repo> <carpeta_salida>"""
import copy, os, subprocess, sys
from docx import Document
REPO, OUT = sys.argv[1], sys.argv[2]
V3 = f"{REPO}/2_Clases/Apuntes/Nuevo_Cajetin_Vistas_Cortes/Apuntes_Unidad1_PCI1119_v3.docx"
M1 = ("Módulo 1 · Geometría Descriptiva", "Modulo_1_Geometria_Descriptiva")
M2 = ("Módulo 2 · Fundamentos del dibujo CAD", "Modulo_2_Fundamentos_del_dibujo_CAD")

def cabeceras(d, modulo, titulo):
    for s in d.sections:
        hp = s.header.paragraphs[0]; r = hp.runs
        r[0].text = "Universidad Católica de Temuco"; r[-1].text = modulo
        p2 = copy.deepcopy(hp._p); hp._p.addnext(p2)
        from docx.text.paragraph import Paragraph
        q = Paragraph(p2, hp._parent); q.runs[0].text = ""; q.runs[-1].text = titulo
        f = s.footer.paragraphs[0].runs
        f[0].text = "Miguel Godoy"; f[2].text = ""

def portada(d, modulo, titulo):
    ps = d.paragraphs
    ps[0].runs[0].text = titulo
    for r in ps[0].runs[1:]: r.text = ""
    ps[1].runs[0].text = f"{modulo} · PCI 1119 Sistemas de Representación"
    for r in ps[1].runs[1:]: r.text = ""

def guardar(d, mod, nombre):
    carpeta = os.path.join(OUT, mod[1]); os.makedirs(carpeta, exist_ok=True)
    ruta = os.path.join(carpeta, nombre + ".docx"); d.save(ruta)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", carpeta, ruta], capture_output=True, check=True)

def capitulo(i, n, nombre):
    d = Document(V3); ps = d.paragraphs
    h1 = [p for p in ps if p.style.name == "Heading 1"]
    titulo = h1[i].text
    ini = h1[i]._p; fin = h1[i + 1]._p if i + 1 < len(h1) else None
    body = d.element.body; keep = {ps[0]._p, ps[1]._p}
    estado = "antes"
    for el in list(body):
        if el is ini: estado = "dentro"; body.remove(el); continue
        if fin is not None and el is fin: estado = "despues"
        if el in keep or el.tag.endswith("sectPr"): continue
        if estado != "dentro": body.remove(el)
    for p in d.paragraphs:  # sube un nivel los encabezados
        if p.style.name == "Heading 2": p.style = d.styles["Heading 1"]
        elif p.style.name == "Heading 3": p.style = d.styles["Heading 2"]
    portada(d, M1[0], titulo); cabeceras(d, M1[0], titulo)
    guardar(d, M1, f"M1_{n:02d}_{nombre}")

def existente(ruta, mod, n, nombre, titulo):
    d = Document(ruta); portada(d, mod[0], titulo); cabeceras(d, mod[0], titulo)
    guardar(d, mod, f"M{mod[0][7]}_{n:02d}_{nombre}")

capitulo(0, 1, "Introduccion_a_los_sistemas_de_representacion")
capitulo(1, 2, "Elementos_del_dibujo")
capitulo(2, 3, "Vistas_de_un_objeto")
existente(f"{REPO}/propuestas/02_Apuntes/Apunte_La_Tercera_Vista.docx", M1, 4, "La_tercera_vista", "La tercera vista")
capitulo(3, 5, "Vistas_auxiliares")
capitulo(4, 6, "Cortes_y_secciones")
existente(f"{REPO}/propuestas/02_Apuntes/Apunte_Isometrico_en_AutoCAD.docx", M2, 1, "Dibujo_isometrico_en_AutoCAD", "Dibujo isométrico en AutoCAD")
